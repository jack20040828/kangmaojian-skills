#!/usr/bin/env python3
"""Generate a formal Word deliverable from reviewer-approved opinions."""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image


SCRIPT_DIR = Path(__file__).resolve().parent
VALIDATOR = SCRIPT_DIR / "validate_delivery_package.py"
NORMALIZER = SCRIPT_DIR / "normalize_opinion_types.py"


def safe_name(value: str) -> str:
    value = re.sub(r"[\\/:*?\"<>|]+", "-", value.strip())
    return value or "未命名"


def set_run_font(run, size: float, bold: bool = False) -> None:
    run.font.name = "宋体"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def add_text(
    doc: Document,
    text: str,
    *,
    size: float = 10.5,
    bold: bool = False,
    align=None,
    before: float = 0,
    after: float = 4,
    keep_next: bool = False,
):
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = 1.18
    fmt.keep_with_next = keep_next
    fmt.keep_together = True
    fmt.widow_control = True
    run = paragraph.add_run(text)
    set_run_font(run, size, bold)
    return paragraph


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_run_font(run, 9)


def strip_label(value: str, label: str) -> str:
    value = value.strip()
    return value[len(label) :].strip() if value.startswith(label) else value


def sentence(value: str) -> str:
    value = value.strip()
    return value if value.endswith(("。", "！", "？", ".", ":", "：")) else value + "。"


def resolve(root: Path, value: str) -> Path:
    path = Path(value)
    return path if path.is_absolute() else root / path


def evidence_images(item: dict) -> list[dict]:
    if "evidence_images" in item:
        value = item.get("evidence_images")
        return value if isinstance(value, list) else []
    screenshot = item.get("screenshot")
    return [screenshot] if isinstance(screenshot, dict) and screenshot else []


def picture_width(path: Path, max_width: float = 6.05, max_height: float = 5.25) -> Inches:
    with Image.open(path) as image:
        width, height = image.size
    if width <= 0 or height <= 0:
        return Inches(max_width)
    return Inches(min(max_width, max_height * width / height))


def add_item(doc: Document, root: Path, item: dict) -> None:
    number = item["item_no"]
    drawing = sentence(item["drawing_refs"])
    add_text(doc, f"{number}、涉及图纸：{drawing}", bold=True, before=4, after=3, keep_next=True)
    opinion = strip_label(item["opinion_text"], "【审查意见】：")
    add_text(doc, f"【审查意见】：{opinion}", after=5, keep_next=True)

    for screenshot in evidence_images(item):
        strategy = str(screenshot.get("strategy", "pdf_provenance")).strip()
        if strategy == "preserve_approved_docx":
            raise ValueError(
                f"Item {number} contains preserved approved-DOCX evidence; "
                "edit the approved base DOCX instead of generating a fresh report"
            )
        image_path = resolve(root, screenshot["image_path"])
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt = paragraph.paragraph_format
        fmt.space_before = Pt(1)
        fmt.space_after = Pt(5)
        fmt.line_spacing = 1.0
        fmt.keep_with_next = True
        fmt.keep_together = True
        paragraph.add_run().add_picture(str(image_path), width=picture_width(image_path))

    regulation = strip_label(item["regulation_text"], "【法规条文】：")
    add_text(doc, f"【法规条文】：{regulation}", after=3, keep_next=True)
    opinion_type = strip_label(item["opinion_type"], "【意见类型】：")
    add_text(doc, f"【意见类型】：{sentence(opinion_type)}", after=7)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.45)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)


def build_report(root: Path, manifest: dict, output: Path) -> Path:
    doc = Document()
    configure_document(doc)

    add_text(doc, manifest["project_name"], size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=8)
    add_text(doc, manifest["report_title"], size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    add_text(doc, manifest["report_date"], size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)

    items_by_section: dict[str, list[dict]] = {}
    for item in manifest["items"]:
        items_by_section.setdefault(item["section"], []).append(item)

    for section_row in manifest["sections"]:
        if isinstance(section_row, dict):
            heading = section_row["heading"].strip()
            empty_text = section_row.get("empty_text", "无意见").strip()
        else:
            heading = str(section_row).strip()
            empty_text = "无意见"
        section_items = items_by_section.get(heading, [])
        heading_text = heading if section_items else f"{heading}{empty_text}"
        add_text(doc, heading_text, size=13, bold=True, before=5, after=6, keep_next=bool(section_items))
        for item in section_items:
            add_item(doc, root, item)

    footer_note = manifest.get("footer_note", "").strip()
    if footer_note:
        add_text(doc, footer_note, size=9.5, before=6, after=0)

    for section in doc.sections:
        section.footer_distance = Inches(0.45)
        footer = section.footer
        if not footer.paragraphs:
            footer.add_paragraph()
        footer.paragraphs[0].clear()
        add_page_number(footer.paragraphs[0])

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workspace", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    root = args.workspace.resolve()
    subprocess.run([sys.executable, str(NORMALIZER), str(root)], check=True)
    subprocess.run([sys.executable, str(VALIDATOR), str(root)], check=True)
    manifest = json.loads((root / "delivery_manifest.json").read_text(encoding="utf-8"))
    default_name = f"【内审意见】{safe_name(manifest['project_name'])}_{safe_name(manifest['report_date'])}.docx"
    output = args.output or root / "output" / default_name
    print(build_report(root, manifest, output))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
