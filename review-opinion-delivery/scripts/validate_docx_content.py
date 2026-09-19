#!/usr/bin/env python3
"""Verify a DOCX against complete opinion blocks in an approved manifest."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


FIELD_LABELS = ["涉及图纸", "【审查意见】", "【法规条文】", "【意见类型】"]
CHINESE_SECTION_PREFIX = re.compile(r"^[一二三四五六七八九十百]+、")
ITEM_START = re.compile(r"^\s*(\d+)\s*[、.]?\s*涉及图纸\s*[：:]")
PROCESS_TRACE_PATTERNS = {
    "笔记": re.compile(r"笔记", re.I),
    "得到大脑": re.compile(r"得到大脑", re.I),
    "note_id": re.compile(r"note[_ -]?id", re.I),
    "证据来源": re.compile(r"证据来源", re.I),
    "PDF页码": re.compile(r"pdf\s*第?\s*\d+\s*页", re.I),
    "标记编号": re.compile(r"标记(?:号|编号)?\s*[：:]?\s*[a-z0-9-]+", re.I),
    "裁剪坐标": re.compile(r"裁剪(?:坐标|范围)|crop[_ -]?(?:box|coords?)", re.I),
    "红框说明": re.compile(r"红框(?:位置|范围|说明|目标)?", re.I),
}


def strip_label(value: str, labels: tuple[str, ...]) -> str:
    text = str(value or "").strip()
    for label in labels:
        if text.startswith(label):
            return text[len(label) :].lstrip("：: ")
    return text


def normalized(value: str) -> str:
    return (
        re.sub(r"\s+", "", str(value or ""))
        .replace(":", "：")
        .replace(",", "，")
        .replace("(", "（")
        .replace(")", "）")
    )


def normalized_heading(value: str) -> str:
    text = normalized(value)
    return CHINESE_SECTION_PREFIX.sub("", text).rstrip("：:")


def body_paragraphs(doc: Document) -> list[Paragraph]:
    """Return body and table paragraphs in document order, including image-only paragraphs."""
    return [Paragraph(element, doc) for element in doc.element.body.iter(qn("w:p"))]


def image_count(paragraph: Paragraph) -> int:
    return len(paragraph._p.xpath(".//a:blip"))


def evidence_count(item: dict) -> int:
    if item.get("image_count") not in (None, ""):
        try:
            return int(item["image_count"])
        except (TypeError, ValueError):
            return -1
    if "evidence_images" in item:
        value = item.get("evidence_images")
        return len(value) if isinstance(value, list) else -1
    return 1 if isinstance(item.get("screenshot"), dict) and item.get("screenshot") else 0


def expected_fields(item: dict) -> list[tuple[str, str]]:
    return [
        ("涉及图纸", strip_label(item.get("drawing_refs", ""), ("涉及图纸：", "涉及图纸:"))),
        ("【审查意见】", strip_label(item.get("opinion_text", ""), ("【审查意见】：", "【审查意见】:"))),
        ("【法规条文】", strip_label(item.get("regulation_text", ""), ("【法规条文】：", "【法规条文】:"))),
        ("【意见类型】", strip_label(item.get("opinion_type", ""), ("【意见类型】：", "【意见类型】:"))),
    ]


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workspace", type=Path)
    parser.add_argument("docx", type=Path)
    args = parser.parse_args()
    root = args.workspace.resolve()
    docx_path = args.docx.resolve()
    errors: list[str] = []

    if not docx_path.exists() or docx_path.stat().st_size == 0:
        print(f"FAIL: DOCX not found or empty: {docx_path}")
        return 1

    manifest = json.loads((root / "delivery_manifest.json").read_text(encoding="utf-8"))
    v13 = manifest.get("schema_version") == "1.3"
    if v13:
        from delivery_contract import validate_contract
        errors.extend(validate_contract(root, manifest))
    doc = Document(docx_path)
    if v13:
        from sync_docx_references import reference_errors
        errors.extend(reference_errors(doc, manifest))
    paragraphs = body_paragraphs(doc)
    texts = [paragraph.text.strip() for paragraph in paragraphs]
    all_body = normalized("\n".join(texts))
    if v13 and manifest.get("footer_note") and normalized(manifest['footer_note']) not in all_body:
        errors.append("Approved final opinion-type note is missing")
    visible_text = "\n".join(
        [
            *texts,
            *(paragraph.text for section in doc.sections for paragraph in section.header.paragraphs),
            *(paragraph.text for section in doc.sections for paragraph in section.footer.paragraphs),
        ]
    )

    for label, pattern in PROCESS_TRACE_PATTERNS.items():
        check_text = visible_text
        if v13 and label == "PDF页码" and manifest.get("location_policy", {}).get("mode") == "pdf_page_title":
            from delivery_contract import PDF_REF
            check_text = PDF_REF.sub("", check_text)
        if pattern.search(check_text):
            errors.append(f"Formal DOCX contains internal process trace: {label}")

    for field in ["project_name", "report_title", "report_date"]:
        expected = normalized(manifest.get(field, ""))
        if expected and expected not in all_body:
            errors.append(f"Document metadata missing {field}: {manifest.get(field)!r}")

    section_positions: list[int] = []
    for section_row in manifest.get("sections", []):
        heading = section_row.get("heading", "") if isinstance(section_row, dict) else str(section_row)
        target = normalized_heading(heading)
        position = next(
            (index for index, text in enumerate(texts) if normalized_heading(text).startswith(target)),
            -1,
        )
        if position < 0:
            errors.append(f"Section heading missing (manual or automatic numbering accepted): {heading!r}")
        section_positions.append(position)
    visible_positions = [value for value in section_positions if value >= 0]
    if visible_positions != sorted(visible_positions):
        errors.append(f"Section headings are out of order: {section_positions}")

    starts: list[tuple[int, int]] = []
    for index, text in enumerate(texts):
        match = ITEM_START.match(text)
        if match:
            starts.append((int(match.group(1)), index))
    expected_numbers = [int(item["item_no"]) for item in manifest.get("items", [])]
    actual_numbers = [number for number, _ in starts]
    if actual_numbers != expected_numbers:
        errors.append(f"Opinion block numbers mismatch: expected {expected_numbers}, found {actual_numbers}")

    start_by_number = {number: index for number, index in starts}
    ordered_start_indexes = [index for _, index in starts]
    for item in manifest.get("items", []):
        number = int(item["item_no"])
        start = start_by_number.get(number)
        if start is None:
            continue
        start_position = ordered_start_indexes.index(start)
        end = ordered_start_indexes[start_position + 1] if start_position + 1 < len(ordered_start_indexes) else len(paragraphs)
        block_text = normalized("\n".join(texts[start:end]))
        if v13:
            from delivery_contract import CONTINUATION, docx_evidence_errors
            from compare_reviewer_docx import FIELD_PATTERNS
            errors.extend(docx_evidence_errors(doc, paragraphs[start:end], item))
            clean_text = "\n".join(t for t in texts[start:end] if not CONTINUATION.fullmatch(t))
            for field, pattern in FIELD_PATTERNS.items():
                match = pattern.search(clean_text)
                if not match or normalized(match.group(1)).rstrip("。") != normalized(item.get(field)).rstrip("。"):
                    errors.append(f"Item {number}: unapproved DOCX field change: {field}")

        field_positions: list[int] = []
        for label, expected_value in expected_fields(item):
            label_position = block_text.find(normalized(label))
            field_positions.append(label_position)
            if label_position < 0:
                errors.append(f"Item {number}: field label missing: {label}")
            if normalized(expected_value) not in block_text:
                errors.append(f"Item {number}: approved {label} content missing: {expected_value!r}")
        present_positions = [position for position in field_positions if position >= 0]
        if present_positions != sorted(present_positions):
            errors.append(f"Item {number}: fields are out of order: {FIELD_LABELS}")

        expected_images = evidence_count(item)
        actual_images = sum(image_count(paragraph) for paragraph in paragraphs[start:end])
        if expected_images < 0:
            errors.append(f"Item {number}: manifest image_count/evidence_images is invalid")
        elif actual_images != expected_images:
            errors.append(f"Item {number}: evidence image count mismatch: expected {expected_images}, found {actual_images}")

    footer_xml = " ".join(section.footer._element.xml for section in doc.sections)
    if not re.search(r"\bPAGE\b", footer_xml):
        errors.append("Footer does not contain a PAGE field")

    if errors:
        for error in errors:
            print(f"FAIL: {error}")
        return 1
    print("PASS: DOCX complete opinion blocks match the approved delivery manifest")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
