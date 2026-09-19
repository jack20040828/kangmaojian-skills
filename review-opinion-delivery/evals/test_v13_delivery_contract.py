"""Anonymous paired positive/negative v1.3 editorial and rendered-document gates."""
from __future__ import annotations

import copy
import csv
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from docx.shared import Inches
from PIL import Image, ImageDraw, ImageFont

SCRIPTS = Path(__file__).resolve().parents[1]/'scripts'
sys.path.insert(0, str(SCRIPTS))
from delivery_contract import validate_contract, docx_evidence_errors
from evidence_geometry import prepare_crop, geometry_errors
from sync_docx_references import synchronize, reference_errors, body_paragraphs
from page_qa_integrity import snapshot, errors as qa_errors, EXTRA_CHECKS
from create_delivery_workspace import VERIFICATION_HEADERS


def run(script, *args, expected=0):
    result = subprocess.run([sys.executable, str(SCRIPTS/script), *map(str,args)], capture_output=True,
                            text=True, encoding='utf-8', env={**os.environ, 'PYTHONUTF8': '1'})
    if result.returncode != expected:
        raise AssertionError(f'{script}: expected {expected}, got {result.returncode}\n{result.stdout}\n{result.stderr}')
    return result.stdout.strip()


def save(root, manifest):
    (root/'delivery_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')


def fixture(parent):
    root = Path(run('create_delivery_workspace.py','匿名整理验收','--root',parent))
    note = root/'source/notes/confirmed.txt'
    note.write_text('匿名审查者确认的地面表达意见；未形成专业技术结论。',encoding='utf-8')
    page = Image.new('RGB',(800,500),'white')
    draw = ImageDraw.Draw(page)
    font_path = Path('C:/Windows/Fonts/arial.ttf')
    font = ImageFont.truetype(str(font_path),22) if font_path.exists() else ImageFont.load_default(size=22)
    draw.rectangle((70,70,730,430),outline='black',width=2)
    draw.line((70,160,730,160),fill='black',width=2)
    draw.text((100,100),'ANONYMOUS GROUND FLOOR DETAIL',font=font,fill='black')
    draw.text((300,245),'GROUND',font=font,fill='black')
    page.save(root/'source/pdf/drawing.pdf','PDF',resolution=144)
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(root/'source/pdf/drawing.pdf')
    rendered_page = pdf[0]
    bitmap = rendered_page.render(scale=2)
    page = bitmap.to_pil().copy()
    bitmap.close();rendered_page.close();pdf.close()
    page.save(root/'renders/source.png')
    # Render-derived coordinates are independent of the fixed problem rectangle.
    image = page.crop((100,100,700,400))
    ImageDraw.Draw(image).rectangle((190,130,310,190),outline='#e32020',width=3)
    image.save(root/'screenshots/ground.png')
    manifest = json.loads((root/'delivery_manifest.json').read_text(encoding='utf-8'))
    item = dict(item_id='OP-010',item_no=1,section='一、设计说明：',source_note_ref='source/notes/confirmed.txt',
                marked_location_ref='source/pdf/drawing.pdf#1',drawing_refs='PDF第1页《匿名构造图》',drawing_no='',
                opinion_text='请核对地面材料表达。',regulation_text='无',opinion_type='设计深度，建议修改（其它）',
                editorial_change_note='无；匿名回归非技术意见',reviewer_confirmed=True,source_ids=['S-01'],
                approval_ref='source/notes/confirmed.txt',image_count=1,
                claims=[dict(claim_id='C-1',text_quote='地面材料',object_label='地面',requires_image=True)],
                evidence_images=[dict(evidence_id='E-1',strategy='pdf_provenance',image_path='screenshots/ground.png',
                                      source_pdf='source/pdf/drawing.pdf',source_page=1,crop_box=[100,100,700,400],
                                      problem_boxes=[[290,230,410,290]],source_size=[800,500],red_box_target='GROUND文字',
                                      context_anchor='GROUND FLOOR表头及构造边界',word_crop=[0,0,0,0],display_width_inches=5.5,
                                      image_sha256=hashlib.sha256((root/'screenshots/ground.png').read_bytes()).hexdigest(),
                                      observed_evidence=[dict(claim_id='C-1',object_label='地面',visible=True,observation='GROUND文字与表头可见')])])
    item['approved_content']={k:item[k] for k in ['opinion_text','regulation_text','opinion_type','drawing_refs']}
    manifest.update(items=[item],location_policy=dict(mode='pdf_page_title',authorization_ref='anonymous user instruction'),
                    source_items=[dict(source_id='S-01',source_ref='source/notes/confirmed.txt',disposition='included',reason='confirmed',item_ids=['OP-010'])])
    save(root,manifest)
    with (root/'verification_log.csv').open('w',encoding='utf-8',newline='') as handle:
        writer=csv.DictWriter(handle,fieldnames=VERIFICATION_HEADERS);writer.writeheader()
        writer.writerow({k: 'OP-010' if k=='item_id' else 1 if k=='item_no' else '匿名固定样例' if k=='notes' else '通过' for k in VERIFICATION_HEADERS})
    run('snapshot_source_integrity.py',root)
    return root,json.loads((root/'delivery_manifest.json').read_text(encoding='utf-8'))


class DeliveryContractTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(prefix='editorial-v13-')
        self.root,self.good=fixture(Path(self.tmp.name))

    def tearDown(self): self.tmp.cleanup()

    def rejected(self, change, expected):
        self.assertEqual(validate_contract(self.root,self.good),[])
        bad=copy.deepcopy(self.good);change(bad)
        self.assertTrue(any(expected in x for x in validate_contract(self.root,bad)),validate_contract(self.root,bad))

    def test_withdrawn_source(self):
        self.rejected(lambda m:m['source_items'][0].update(disposition='withdrawn',confirmation_ref='author',item_ids=[]),'withdrawn')
        self.good['source_items'].append(dict(source_id='S-02',source_ref='record',disposition='withdrawn',confirmation_ref='author',reason='withdrawn by author',item_ids=[]))
        self.assertEqual(validate_contract(self.root,self.good),[])

    def test_source_conflict(self):
        self.rejected(lambda m:m.update(source_conflicts=[dict(status='pending')]),'conflict')
        self.good['source_conflicts']=[dict(status='resolved',resolution_ref='author confirmation')]
        self.assertEqual(validate_contract(self.root,self.good),[])

    def test_unauthorized_wording(self):
        self.rejected(lambda m:m['items'][0].update(opinion_text='请核对地面材料表达，并修改全部立面。'),'unapproved change')

    def test_reviewer_dimension_not_law(self):
        def mutate(m):
            item=m['items'][0];item['value_origins']=[dict(value='1.05m',origin='reviewer_remedy',source_ref='author')]
            item['regulation_text']=item['approved_content']['regulation_text']='洞口最低1.05m'
        self.rejected(mutate,'code minimum')
        self.good['items'][0]['value_origins']=[dict(value='1.05m',origin='reviewer_remedy',source_ref='author')]
        self.assertEqual(validate_contract(self.root,self.good),[])
        self.good['items'][0]['opinion_text']=self.good['items'][0]['approved_content']['opinion_text']='请核对地面材料，调整厚度20mm。'
        self.assertTrue(any('origins' in e for e in validate_contract(self.root,self.good)))
        self.good['items'][0]['value_origins'].append(dict(value='20mm',origin='reviewer_remedy',source_ref='author'))
        self.assertEqual(validate_contract(self.root,self.good),[])

    def test_missing_decisive_sentence(self):
        self.rejected(lambda m:m['items'][0]['evidence_images'][0]['observed_evidence'][0].update(visible=False),'not visibly')

    def test_floor_object_replacement(self):
        self.rejected(lambda m:m['items'][0]['evidence_images'][0]['observed_evidence'][0].update(object_label='楼面'),'object differs')

    def test_scope_requires_confirmation(self):
        self.rejected(lambda m:m['items'][0].update(scope_changes=[dict(included=True,status='pending')]),'scope expansion')
        self.good['items'][0]['scope_changes']=[dict(included=True,status='approved',approval_ref='author')]
        self.assertEqual(validate_contract(self.root,self.good),[])

    def test_pdf_locator_exception(self):
        self.rejected(lambda m:m['location_policy'].update(authorization_ref=''),'authorization')
        self.rejected(lambda m:m['items'][0].update(drawing_refs='PDF第0页《图》'),'locator must')
        self.rejected(lambda m:m['items'][0].update(drawing_no='invented'),'locator must')
        item=self.good['items'][0];item['drawing_refs']=item['approved_content']['drawing_refs']='A-01 匿名图'
        self.good['location_policy']={'mode':'drawing_number_title'}
        self.assertEqual(validate_contract(self.root,self.good),[])

    def test_image_history(self):
        self.rejected(lambda m:m.update(evidence_history=[dict(evidence_id='E-1',source_ref='previous DOCX',reason='duplicate',disposition='removed')]),'removed image')
        self.good['evidence_history']=[dict(evidence_id='E-old',source_ref='previous DOCX',reason='clearer same object',disposition='replaced',replacement_ids=['E-1'])]
        self.assertEqual(validate_contract(self.root,self.good),[])

    def test_preserved_docx_not_exempt(self):
        image=self.good['items'][0]['evidence_images'][0];image['strategy']='preserve_approved_docx'
        self.assertEqual(validate_contract(self.root,self.good),[])
        self.rejected(lambda m:m['items'][0]['evidence_images'][0].update(observed_evidence=[]),'not exempt')

    def test_crop_geometry_and_color(self):
        boxes=[[100,100,140,120],[160,100,200,120]]
        original=copy.deepcopy(boxes)
        crop,limits=prepare_crop([90,50,210,170],boxes,[400,300])
        self.assertEqual(crop,[25,50,275,170]);self.assertEqual(boxes,original);self.assertEqual(limits,[])
        e=dict(crop_box=crop,problem_boxes=boxes,source_size=[400,300],context_anchor='axis')
        self.assertEqual(geometry_errors(e),[])
        e['crop_box']=[90,50,210,170];self.assertTrue(geometry_errors(e))
        e.update(crop_box=crop,mark_color='yellow');self.assertTrue(geometry_errors(e))
        e['color_authorization_ref']='user';self.assertEqual(geometry_errors(e),[])
        e.update(crop_box=[0,0,100,100],problem_boxes=[[0,0,80,80]],source_size=[100,100])
        self.assertTrue(geometry_errors(e));e['boundary_limitation']='small sheet; use detail companion';self.assertEqual(geometry_errors(e),[])

    def test_evidence_hash_width_crop(self):
        item=self.good['items'][0];doc=Document();p=doc.add_paragraph()
        p.add_run().add_picture(str(self.root/'screenshots/ground.png'),width=Inches(5.5))
        self.assertEqual(docx_evidence_errors(doc,[p],item),[])
        for key,value,expected in [('image_sha256','0'*64,'hash/order'),('display_width_inches',3,'display size'),('word_crop',[0.1,0,0,0],'crop')]:
            bad=copy.deepcopy(item);bad['evidence_images'][0][key]=value
            self.assertTrue(any(expected in e for e in docx_evidence_errors(doc,[p],bad)))

    def test_stable_reference_updates(self):
        doc=Document();p=doc.add_paragraph();p.add_run('第');p.add_run('9');p.add_run('条续')
        self.good['docx_references']=[dict(paragraph_index=0,occurrence=0,item_id='OP-010')]
        self.assertTrue(reference_errors(doc,self.good))
        self.assertTrue(docx_evidence_errors(doc,[p],self.good['items'][0]))
        synchronize(doc,self.good,[dict(paragraph_index=0,before_text='第9条续',kind='continuation',item_id='OP-010')])
        self.assertEqual(p.text,'第1条续');self.assertEqual(len(p.runs),3)
        self.assertEqual(reference_errors(doc,self.good),[])
        doc.add_paragraph('【法规条文】：示例条例第12条。')
        self.assertEqual(reference_errors(doc,self.good),[])
        with self.assertRaises(ValueError):
            synchronize(doc,self.good,[dict(paragraph_index=0,before_text='wrong anchor',kind='continuation',item_id='OP-010')])

    def test_composer_fixed_problem_coordinates(self):
        from compose_evidence_cards import build_panel
        spec=dict(source_pdf='source/pdf/drawing.pdf',source_page=1,source_image='renders/source.png',
                  crop=[280,210,440,310],source_problem_boxes=[[290,230,410,290]],mark_color='yellow',color_authorization_ref='user')
        panel=build_panel(self.root,spec,300)
        self.assertEqual(spec['resolved_geometry']['problem_boxes'],[[290,230,410,290]])
        self.assertEqual(spec['resolved_geometry']['crop_box'],[200,185,500,335])
        self.assertEqual(panel.getpixel((90,45)),(255,255,0))
        del spec['color_authorization_ref']
        with self.assertRaises(ValueError): build_panel(self.root,spec,300)

    def test_full_package_docx_and_mutation(self):
        run('validate_delivery_package.py',self.root)
        run('generate_delivery_report.py',self.root)
        report=next((self.root/'output').glob('*.docx'))
        run('validate_docx_content.py',self.root,report)
        doc=Document(report)
        next(p for p in doc.paragraphs if p.text.startswith('【法规条文】')).insert_paragraph_before('第9条续')
        doc.save(self.root/'output/stale-continuation.docx')
        self.assertIn('stale',run('validate_docx_content.py',self.root,self.root/'output/stale-continuation.docx',expected=1))
        doc=Document(report);doc.add_paragraph('笔记编号：匿名过程记录')
        doc.save(self.root/'output/process-trace.docx')
        self.assertIn('process trace',run('validate_docx_content.py',self.root,self.root/'output/process-trace.docx',expected=1))
        doc=Document(report)
        next(p for p in doc.paragraphs if p.text.startswith('【审查意见】')).add_run('同时修改全部立面。')
        doc.save(self.root/'output/tampered.docx')
        self.assertIn('unapproved DOCX field',run('validate_docx_content.py',self.root,self.root/'output/tampered.docx',expected=1))
        self.good['source_conflicts']=[dict(status='pending')];save(self.root,self.good)
        run('generate_delivery_report.py',self.root,expected=1)

    def test_normalization_and_legacy_no_migration(self):
        item=self.good['items'][0];item['opinion_type']=item['approved_content']['opinion_type']='设计深度文，建议修改（其它）'
        save(self.root,self.good);run('normalize_opinion_types.py',self.root);run('validate_delivery_package.py',self.root)
        parent=self.root.parent
        before=(self.root/'delivery_manifest.json').read_bytes()
        run('create_delivery_workspace.py','匿名整理验收','--root',parent,'--schema-version','1.2',expected=1)
        self.assertEqual(before,(self.root/'delivery_manifest.json').read_bytes())

    def test_page_qa_complete_and_fresh(self):
        report=self.root/'output/sample.docx';Document().save(report)
        pdf=self.root/'source/pdf/drawing.pdf';render=self.root/'renders/source.png'
        qa=self.root/'report_qa.csv';snapshot(qa,report,pdf,[render])
        rows=[dict(report_docx=str(report),report_pdf=str(pdf),render_path=str(render),page_no=1,**{k:'通过' for k in EXTRA_CHECKS})]
        self.assertEqual(qa_errors(qa,rows),[])
        self.assertTrue(qa_errors(qa,rows*2))
        rows[0]['image_crop_check']='待检查';self.assertTrue(qa_errors(qa,rows));rows[0]['image_crop_check']='通过'
        self.good['report_date']='changed';save(self.root,self.good)
        self.assertTrue(any('stale' in e for e in qa_errors(qa,rows)))


if __name__ == '__main__':
    if len(sys.argv)>2 and sys.argv[1]=='--emit-fixture':
        root,manifest=fixture(Path(sys.argv[2]));run('generate_delivery_report.py',root)
        run('validate_docx_content.py',root,next((root/'output').glob('*.docx')));print(root)
    else:
        unittest.main()
