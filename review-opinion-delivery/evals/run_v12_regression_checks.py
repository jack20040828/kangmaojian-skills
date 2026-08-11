#!/usr/bin/env python3
"""Run v1.2 correction, block-layout, evidence, and semantic-diff regressions."""

from __future__ import annotations

import copy
import csv
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image


SKILL_ROOT = Path(__file__).resolve().parent.parent
SCRIPTS = SKILL_ROOT / "scripts"
CHECK_FIELDS = [
    "note_transcription_check",
    "author_intent_check",
    "drawing_reference_check",
    "numeric_scope_check",
    "marked_location_check",
    "regulation_check",
    "opinion_type_check",
    "editorial_change_check",
]


def run(script: str, *args: object, expected: int = 0) -> subprocess.CompletedProcess[str]:
    result = subprocess.run(
        [sys.executable, str(SCRIPTS / script), *(str(arg) for arg in args)],
        check=False,
        capture_output=True,
        text=True,
        encoding="utf-8",
        env={**os.environ, "PYTHONIOENCODING": "utf-8"},
    )
    if result.returncode != expected:
        raise AssertionError(
            f"{script} returned {result.returncode}, expected {expected}\nSTDOUT:\n{result.stdout}\nSTDERR:\n{result.stderr}"
        )
    return result


def write_json(path: Path, value: dict) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_verification(path: Path, items: list[dict]) -> None:
    headers = [
        "item_id",
        "item_no",
        *CHECK_FIELDS[:5],
        "pdf_provenance_check",
        *CHECK_FIELDS[5:],
        "result",
        "notes",
    ]
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        for item in items:
            pdf_evidence = any(
                image.get("strategy", "pdf_provenance") == "pdf_provenance"
                for image in item.get("evidence_images", [])
            )
            row = {
                "item_id": item["item_id"],
                "item_no": item["item_no"],
                "pdf_provenance_check": "通过" if pdf_evidence else "不适用",
                "result": "通过",
                "notes": "匿名v1.2回归样例",
            }
            row.update({field: "通过" for field in CHECK_FIELDS})
            writer.writerow(row)


def add_page_field(doc: Document) -> None:
    paragraph = doc.sections[0].footer.paragraphs[0]
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def build_block_docx(path: Path, manifest: dict, image: Path) -> None:
    doc = Document()
    doc.add_paragraph(manifest["project_name"])
    doc.add_paragraph(manifest["report_title"])
    doc.add_paragraph(manifest["report_date"])
    sections = ["设计说明：", "平面图：", "立面、剖面图：", "大样图：无意见"]
    by_section = {row["heading"]: [] for row in manifest["sections"]}
    for item in manifest["items"]:
        by_section[item["section"]].append(item)
    for index, section_row in enumerate(manifest["sections"]):
        doc.add_paragraph(sections[index])
        for item in by_section[section_row["heading"]]:
            text = (
                f"{item['item_no']}、涉及图纸：{item['drawing_refs']} "
                f"【审查意见】：{item['opinion_text']} "
                f"【法规条文】：{item['regulation_text']} "
                f"【意见类型】：{item['opinion_type']}"
            )
            doc.add_paragraph(text)
            for _ in item.get("evidence_images", []):
                doc.add_paragraph().add_run().add_picture(str(image))
    add_page_field(doc)
    doc.save(path)


def build_numbered_docx(path: Path, count: int, regulation_prefix: str) -> None:
    doc = Document()
    doc.add_paragraph("匿名项目")
    doc.add_paragraph("建筑施工图内审意见")
    doc.add_paragraph("2026年8月5日")
    doc.add_paragraph("一、设计说明：")
    for number in range(1, count + 1):
        regulation = "无" if regulation_prefix == "无" else f"{regulation_prefix}{number}"
        doc.add_paragraph(
            f"{number}、涉及图纸：A-{number:02d} 匿名图 "
            f"【审查意见】：匿名意见{number} "
            f"【法规条文】：{regulation} "
            "【意见类型】：设计深度，必须修改（其它）"
        )
    add_page_field(doc)
    doc.save(path)


def build_semantic_docx(path: Path, items: list[dict], image: Path) -> None:
    doc = Document()
    doc.add_paragraph("匿名项目")
    doc.add_paragraph("建筑单体施工图内审意见")
    doc.add_paragraph("2026年8月9日")
    doc.add_paragraph("设计说明：")
    for item in items:
        doc.add_paragraph(
            f"{item['item_no']}、涉及图纸：{item['drawing_refs']} "
            f"【审查意见】：{item['opinion_text']} "
            f"【法规条文】：{item.get('regulation_text', '无')} "
            f"【意见类型】：{item.get('opinion_type', '设计深度，必须修改（其它）')}"
        )
        for _ in range(item.get("image_count", 0)):
            doc.add_paragraph().add_run().add_picture(str(image))
    add_page_field(doc)
    doc.save(path)


def main() -> int:
    with tempfile.TemporaryDirectory(prefix="review-opinion-v12-") as temporary:
        root = Path(temporary)
        created = run("create_delivery_workspace.py", "匿名v1.2项目", "--root", root)
        workspace = Path(created.stdout.strip().splitlines()[-1])
        source_pdf = workspace / "source" / "pdf" / "marked.pdf"
        source_pdf.write_bytes(b"%PDF-1.4\nanonymous fixture\n%%EOF")
        source_docx = workspace / "source" / "docx" / "approved.docx"
        Document().save(source_docx)
        (workspace / "source" / "notes" / "reviewer.txt").write_text("匿名审查人意见", encoding="utf-8")
        evidence = workspace / "screenshots" / "evidence.png"
        Image.new("RGB", (120, 80), "white").save(evidence)
        run("snapshot_source_integrity.py", workspace)

        pdf_evidence = {
            "strategy": "pdf_provenance",
            "image_path": "screenshots/evidence.png",
            "source_pdf": "source/pdf/marked.pdf",
            "source_page": 1,
            "crop_box": [0, 0, 100, 100],
            "red_box_target": "匿名位置",
        }
        manifest_path = workspace / "delivery_manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        manifest["items"] = [
            {
                "item_id": "OP-001",
                "item_no": 1,
                "section": "一、设计说明：",
                "source_note_ref": "笔记1",
                "marked_location_ref": "PDF第1页",
                "drawing_refs": "SM-01 建筑说明",
                "opinion_text": "确定性类型错字必须改正。",
                "regulation_text": "无",
                "opinion_type": "设计深度文，必须修改（其它）",
                "editorial_change_note": "无",
                "reviewer_confirmed": True,
                "image_count": 2,
                "evidence_images": [copy.deepcopy(pdf_evidence), copy.deepcopy(pdf_evidence)],
            },
            {
                "item_id": "OP-002",
                "item_no": 2,
                "section": "二、平面图：",
                "source_note_ref": "笔记2",
                "marked_location_ref": "文字意见",
                "drawing_refs": "JS-01 一层平面图",
                "opinion_text": "必须状态不得被纠错程序改变。",
                "regulation_text": "无",
                "opinion_type": "一般性条文，必须修改（其它）",
                "editorial_change_note": "无",
                "reviewer_confirmed": True,
                "image_count": 0,
                "evidence_images": [],
            },
            {
                "item_id": "OP-003",
                "item_no": 3,
                "section": "三、立面、剖面图：",
                "source_note_ref": "笔记3",
                "marked_location_ref": "最终Word原图",
                "drawing_refs": "JS-05 剖面图",
                "opinion_text": "建议状态和最终Word证据均须保留。",
                "regulation_text": "无",
                "opinion_type": "设计深度，建议修改（其它）",
                "editorial_change_note": "无",
                "reviewer_confirmed": True,
                "image_count": 1,
                "evidence_images": [
                    {
                        "strategy": "preserve_approved_docx",
                        "source_docx": "source/docx/approved.docx",
                        "source_item_ref": "最终第3条",
                    }
                ],
            },
        ]
        write_json(manifest_path, manifest)
        write_verification(workspace / "verification_log.csv", manifest["items"])

        run("normalize_opinion_types.py", workspace)
        normalized = json.loads(manifest_path.read_text(encoding="utf-8"))
        types = [item["opinion_type"] for item in normalized["items"]]
        if types != [
            "设计深度，必须修改（其它）",
            "一般性条文，必须修改（其它）",
            "设计深度，建议修改（其它）",
        ]:
            raise AssertionError(f"Mandatory/suggested status changed unexpectedly: {types}")
        if any(item["regulation_text"] != "无" for item in normalized["items"]):
            raise AssertionError("Regulation value 无 was repopulated")
        edit_rows = list(csv.DictReader((workspace / "edit_log.csv").open(encoding="utf-8", newline="")))
        if len(edit_rows) != 1 or edit_rows[0]["before"] != "设计深度文，必须修改（其它）":
            raise AssertionError(f"Deterministic correction was not logged once: {edit_rows}")
        run("validate_delivery_package.py", workspace)

        block_docx = workspace / "output" / "combined-fields.docx"
        build_block_docx(block_docx, normalized, evidence)
        run("validate_docx_content.py", workspace, block_docx)

        traced_docx = workspace / "output" / "process-trace.docx"
        traced = Document(block_docx)
        traced.add_paragraph("证据来源：引用笔记 note_id=123，PDF第1页红框位置。")
        traced.save(traced_docx)
        run("validate_docx_content.py", workspace, traced_docx, expected=1)

        card_manifest = workspace / "compact-card.json"
        compact_output = workspace / "screenshots" / "compact.png"
        write_json(
            card_manifest,
            {
                "width": 200,
                "margin": 10,
                "output": str(compact_output),
                "panels": [
                    {
                        "source_pdf": str(source_pdf),
                        "source_page": 1,
                        "source_image": str(evidence),
                        "crop": [0, 0, 120, 80],
                        "red_boxes": [],
                    }
                ],
            },
        )
        run("compose_evidence_cards.py", card_manifest)
        with Image.open(compact_output) as compact:
            if compact.size != (200, 140):
                raise AssertionError(f"Compact evidence default retained decorative header/footer: {compact.size}")

        mixed = copy.deepcopy(manifest)
        mixed["items"][1]["opinion_type"] = "设计深度，必须或建议修改（其它）"
        write_json(manifest_path, mixed)
        run("normalize_opinion_types.py", workspace, expected=1)
        after_failure = json.loads(manifest_path.read_text(encoding="utf-8"))
        if after_failure["items"][0]["opinion_type"] != "设计深度，必须修改（其它）":
            raise AssertionError("Known alias was not corrected before mixed-manifest blocking")
        if after_failure["items"][1]["opinion_type"] != "设计深度，必须或建议修改（其它）":
            raise AssertionError("Ambiguous opinion type was modified instead of blocked")
        mixed_edit_rows = list(csv.DictReader((workspace / "edit_log.csv").open(encoding="utf-8", newline="")))
        if len(mixed_edit_rows) != 2 or mixed_edit_rows[-1]["before"] != "设计深度文，必须修改（其它）":
            raise AssertionError(f"Known correction in mixed manifest was not logged: {mixed_edit_rows}")
        write_json(manifest_path, normalized)

        draft_docx = root / "draft-24.docx"
        final_docx = root / "final-18.docx"
        build_numbered_docx(draft_docx, 24, "旧法规")
        build_numbered_docx(final_docx, 18, "无")
        mapping = root / "mapping.json"
        write_json(mapping, {"final_to_draft": {str(number): number + 6 for number in range(1, 19)}})
        json_output = root / "diff.json"
        markdown_output = root / "diff.md"
        run(
            "compare_reviewer_docx.py",
            draft_docx,
            final_docx,
            "--mapping",
            mapping,
            "--json-output",
            json_output,
            "--markdown-output",
            markdown_output,
        )
        diff = json.loads(json_output.read_text(encoding="utf-8"))
        if diff["summary"]["draft_items"] != 24 or diff["summary"]["final_items"] != 18:
            raise AssertionError(f"24-to-18 semantic diff failed: {diff['summary']}")
        if diff["summary"]["deleted_draft_items"] != [1, 2, 3, 4, 5, 6]:
            raise AssertionError(f"Deleted draft items re-entered mapping: {diff['summary']}")
        if len(diff["regulation_changes"]) != 18 or any(row["after"] != "无" for row in diff["regulation_changes"]):
            raise AssertionError("Regulations changed to 无 were not preserved by semantic diff")

        draft_items = [
            {
                "item_no": number,
                "drawing_refs": f"A-{number:02d} 图纸甲、A-{number:02d}B 图纸乙",
                "opinion_text": f"匿名核心问题{number}，核对图纸范围甲、范围乙和范围丙并统一修改。",
                "image_count": 2 if number <= 14 else 1,
            }
            for number in range(1, 26)
        ]
        primary_drafts = [1, 3, 4, 5, 6, 7, 9, 10, 11, 12, 13, 14, 15, 16, 18, 19, 20, 21, 22, 23, 25]
        final_items = []
        for final_no, draft_no in enumerate(primary_drafts, start=1):
            source = draft_items[draft_no - 1]
            item = dict(source, item_no=final_no, image_count=2 if final_no <= 7 else 1)
            if final_no in {1, 2}:
                item["drawing_refs"] = f"A-{draft_no:02d} 图纸甲"
                item["opinion_text"] = f"匿名核心问题{draft_no}，核对图纸范围甲并修改。"
            if final_no == 8:
                item["opinion_text"] = "匿名核心问题10和17属于同一整改动作，合并修改。"
            if final_no == 14:
                item["opinion_text"] = "匿名核心问题16和24属于同一整改动作，合并修改。"
            final_items.append(item)

        draft_25 = root / "draft-25.docx"
        final_21 = root / "final-21.docx"
        build_semantic_docx(draft_25, draft_items, evidence)
        build_semantic_docx(final_21, final_items, evidence)
        mapping_25_21 = root / "mapping-25-21.json"
        write_json(
            mapping_25_21,
            {
                "final_to_draft": {str(number): draft_no for number, draft_no in enumerate(primary_drafts, start=1)},
                "absorbed": {"8": [17], "14": [24]},
            },
        )
        draft_render = root / "draft-render"
        final_render = root / "final-render"
        draft_render.mkdir()
        final_render.mkdir()
        image_bytes = evidence.read_bytes()
        for number in range(1, 35):
            (draft_render / f"page-{number}.png").write_bytes(image_bytes)
        for number in range(1, 13):
            (final_render / f"page-{number}.png").write_bytes(image_bytes)
        diff_25_21_json = root / "diff-25-21.json"
        diff_25_21_md = root / "diff-25-21.md"
        run(
            "compare_reviewer_docx.py",
            draft_25,
            final_21,
            "--mapping",
            mapping_25_21,
            "--draft-render-dir",
            draft_render,
            "--final-render-dir",
            final_render,
            "--json-output",
            diff_25_21_json,
            "--markdown-output",
            diff_25_21_md,
        )
        diff_25_21 = json.loads(diff_25_21_json.read_text(encoding="utf-8"))
        summary = diff_25_21["summary"]
        expected_summary = {
            "draft_items": 25,
            "final_items": 21,
            "draft_images": 39,
            "final_images": 28,
            "deleted_draft_items": [2, 8],
            "added_final_items": [],
        }
        for key, value in expected_summary.items():
            if summary[key] != value:
                raise AssertionError(f"25-to-21 summary mismatch for {key}: {summary}")
        if diff_25_21["merges"] != [
            {"final_item": 8, "draft_items": [10, 17]},
            {"final_item": 14, "draft_items": [16, 24]},
        ]:
            raise AssertionError(f"25-to-21 merge mapping failed: {diff_25_21['merges']}")
        if [row["final_item"] for row in diff_25_21["narrowed"]] != [1, 2]:
            raise AssertionError(f"Narrowed items were not detected: {diff_25_21['narrowed']}")
        if diff_25_21["documents"]["draft"]["page_count"] != 34 or diff_25_21["documents"]["final"]["page_count"] != 12:
            raise AssertionError("Rendered page counts were not recorded")
        report_text = diff_25_21_md.read_text(encoding="utf-8")
        if "None页" in report_text or "图片由39张减少至28张" not in report_text:
            raise AssertionError("Difference report page/image direction text is incorrect")

        auto_json = root / "auto-diff.json"
        auto_md = root / "auto-diff.md"
        run(
            "compare_reviewer_docx.py",
            draft_25,
            final_21,
            "--json-output",
            auto_json,
            "--markdown-output",
            auto_md,
        )
        auto = json.loads(auto_json.read_text(encoding="utf-8"))
        if auto["summary"]["deleted_draft_items"] or auto["summary"]["added_final_items"]:
            raise AssertionError("Automatic mapping asserted ambiguous deletions or additions")
        if auto["summary"]["mapping_mode"] != "automatic" or not auto["mapping_required"]:
            raise AssertionError("Automatic mapping did not expose unresolved relations for confirmation")

    print("PASS: v1.2 correction, process-trace, compact evidence, type gating, and semantic-diff regressions")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
