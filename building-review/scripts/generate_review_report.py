#!/usr/bin/env python3
"""Generate a validated template-style Word report from building-review issues."""

from __future__ import annotations

import argparse
import json
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from validate_review_package import image_size, resolve_paths, rows, validate_workspace

SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_DIR = SCRIPT_DIR.parent
SINGLE_TEMPLATE = SKILL_DIR / "范本_建筑单体施工图审查意见.docx"
SITE_TEMPLATE = SKILL_DIR / "范本_建筑总图施工图审查意见.docx"

SINGLE_SECTIONS_V11 = [
    "一、设计说明：",
    "二、平面图：",
    "三、立面、剖面图：",
    "四、大样图：",
]
SINGLE_SECTIONS_V12 = [
    "设计说明：",
    "二、平面图：",
    "三、立面、剖面图：",
    "四、大样图：",
]
SITE_SECTIONS = ["总图（建筑）施工图设计说明：", "总图设计图纸："]
SECTION_KEYS = {
    "single": {
        "设计说明": 0,
        "平面图": 1,
        "立面剖面图": 2,
        "大样图": 3,
    },
    "site": {
        "总图设计说明": SITE_SECTIONS[0],
        "总图设计图纸": SITE_SECTIONS[1],
    },
}


def clear_document_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)


def set_default_fonts(doc: Document) -> None:
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        if style_name not in doc.styles:
            continue
        style = doc.styles[style_name]
        style.font.name = "宋体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        if style_name == "Normal":
            style.font.size = Pt(11)


def add_paragraph(
    doc: Document,
    text: str = "",
    *,
    bold: bool = False,
    size: int | None = None,
    align=None,
) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    if align is not None:
        paragraph.alignment = align


def finish_sentence(text: str) -> str:
    text = text.rstrip()
    return text if text.endswith(("。", "！", "？", ".", "!", "?")) else text + "。"


def chinese_date(value: date) -> str:
    return f"{value.year}年{value.month}月{value.day}日"


def safe_filename_component(value: str) -> str:
    value = re.sub(r'[\\/:*?"<>|]+', "-", value.strip())
    return value or "未命名项目"


def classify_single(issue: dict[str, str]) -> str:
    text = " ".join([issue.get("drawing_refs", ""), issue.get("problem", ""), issue.get("screenshot_location", "")])
    if any(key in text for key in ["立面", "剖面"]):
        return SINGLE_SECTIONS_V11[2]
    if any(key in text for key in ["大样", "详图", "节点", "楼梯", "卫生间", "门窗", "墙身", "栏杆"]):
        return SINGLE_SECTIONS_V11[3]
    if "平面" in text:
        return SINGLE_SECTIONS_V11[1]
    return SINGLE_SECTIONS_V11[0]


def classify_site(issue: dict[str, str]) -> str:
    text = " ".join([issue.get("drawing_refs", ""), issue.get("problem", "")])
    return SITE_SECTIONS[0] if "说明" in text else SITE_SECTIONS[1]


def fit_picture_width(path: Path, max_width: float = 6.1, max_height: float = 5.2) -> Inches:
    size = image_size(path)
    if not size:
        return Inches(max_width)
    width, height = size
    if width <= 0 or height <= 0:
        return Inches(max_width)
    return Inches(min(max_width, max_height * width / height))


def add_issue(
    doc: Document,
    root: Path,
    issue: dict[str, str],
    number: int,
    schema_version: str,
) -> None:
    add_paragraph(doc, finish_sentence(f"{number}、涉及图纸：{issue.get('drawing_refs', '').strip()}"), bold=True)
    add_paragraph(doc, f"【审查意见】：{issue.get('problem', '').strip()}")
    for screenshot in resolve_paths(root, issue.get("screenshot_path", "")):
        if screenshot.exists():
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(screenshot), width=fit_picture_width(screenshot))
    if schema_version in {"1.2", "1.3", "1.4"} and issue.get("citation_mode", "").strip().lower() == "none":
        law = "【法规条文】：无"
    else:
        article = issue.get("standard_article", "").strip()
        requirement = issue.get("standard_requirement", "").strip()
        if schema_version in {"1.2", "1.3", "1.4"}:
            source = issue.get("standard_display_name", "").strip()
        else:
            source_value = issue.get("standard_source", "").strip()
            source = Path(source_value).name if source_value else ""
        law = f"【法规条文】：{source}"
        if article:
            law += article
        if requirement:
            law += f"：{requirement}"
    add_paragraph(doc, finish_sentence(law))
    add_paragraph(doc, finish_sentence(f"【意见类型】：{issue.get('opinion_type', '').strip()}"))
    add_paragraph(doc, "")


def load_manifest(root: Path) -> tuple[str, dict]:
    path = root / "review_manifest.json"
    if not path.exists():
        return "", {}
    manifest = json.loads(path.read_text(encoding="utf-8"))
    version = str(manifest.get("schema_version", ""))
    return version if version in {"1.1", "1.2", "1.3", "1.4"} else "", manifest


def build_report(args: argparse.Namespace) -> Path:
    root = args.project_workspace.resolve()
    errors, _warnings = validate_workspace(root)
    if errors:
        raise RuntimeError("review package validation failed:\n- " + "\n- ".join(errors))
    schema_version, manifest = load_manifest(root)
    versioned = schema_version in {"1.1", "1.2", "1.3", "1.4"}
    manifest_type = manifest.get("report_type") if versioned else None
    if args.report_type and manifest_type and args.report_type != manifest_type:
        raise RuntimeError(f"--report-type {args.report_type} conflicts with manifest report_type {manifest_type}")
    report_type = args.report_type or manifest_type or "single"
    template = SINGLE_TEMPLATE if report_type == "single" else SITE_TEMPLATE
    if report_type == "single":
        sections = SINGLE_SECTIONS_V12 if schema_version in {"1.2", "1.3", "1.4"} else SINGLE_SECTIONS_V11
    else:
        sections = SITE_SECTIONS
    classifier = classify_single if report_type == "single" else classify_site
    if report_type == "single" and schema_version in {"1.2", "1.3", "1.4"}:
        default_name = f"【建单内审】{safe_filename_component(args.project_name)}{date.today().isoformat()}.docx"
    else:
        default_name = "建筑单体施工图审查意见.docx" if report_type == "single" else "建筑总图施工图审查意见.docx"
    output = args.output or root / "output" / default_name

    issues = [row for row in rows(root / "issue_candidates.csv") if row.get("status", "").strip() == "verified"]
    if versioned:
        issues.sort(key=lambda row: int(row["display_order"]))
    grouped = {section: [] for section in sections}
    for issue in issues:
        if versioned:
            section_key = SECTION_KEYS[report_type][issue["report_section"].strip()]
            section = sections[section_key] if report_type == "single" else section_key
        else:
            section = classifier(issue)
        grouped[section].append(issue)

    doc = Document(template)
    clear_document_body(doc)
    set_default_fonts(doc)
    add_paragraph(doc, args.project_name, bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
    default_title = "建筑单体施工图内审意见" if report_type == "single" and schema_version in {"1.2", "1.3", "1.4"} else (
        "建筑单体施工图审查意见" if report_type == "single" else "建筑总图施工图审查意见"
    )
    add_paragraph(doc, args.report_title or default_title, bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, args.report_date or chinese_date(date.today()), size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_paragraph(doc, "")
    if not issues:
        add_paragraph(doc, "经审查，本次未形成正式审查意见。", bold=True)
        add_paragraph(doc, "")

    number = 1
    for section in sections:
        add_paragraph(doc, section, bold=True, size=13)
        for issue in grouped[section]:
            add_issue(doc, root, issue, number, schema_version)
            number += 1
        if not grouped[section]:
            add_paragraph(doc, "")
    add_paragraph(doc, "注：意见类型按消防安全强制性条文、一般性条文、政策规定、设计深度及其它强制性/一般性条文等12类填写。")

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("project_workspace", type=Path)
    parser.add_argument("--report-type", choices=["single", "site"])
    parser.add_argument("--project-name", required=True)
    parser.add_argument("--report-title")
    parser.add_argument("--report-date", default="")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    try:
        print(build_report(args))
    except RuntimeError as error:
        print(f"FAIL: {error}")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
