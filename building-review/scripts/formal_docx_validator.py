#!/usr/bin/env python3
"""Validate formal AI-initial DOCX content, template fidelity, and package hygiene."""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from formal_review_report import (
    FOOTER_NOTE,
    PAGINATION_CONTROL_TAGS,
    SINGLE_SECTIONS,
    SINGLE_TEMPLATE,
    SITE_SECTIONS,
    SITE_TEMPLATE,
    finish_sentence,
    regulation_text,
    report_identity,
)
from validate_review_package import resolve_paths, rows, validate_workspace

FORBIDDEN_PROCESS_TEXT = [
    "项目画像",
    "审查边界",
    "风险等级",
    "待人工复核",
    "需人工复核",
    "未经人工",
    "内部日志",
    "截图来源",
    "截图坐标",
    "红框说明",
    "红框坐标",
    "ai_ready",
    "review_stage",
    "gate_origin",
    "AI初审完成",
]


def document_text(doc: Document) -> str:
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def section_metrics(section) -> tuple[int, ...]:
    return tuple(
        int(value)
        for value in [
            section.page_width,
            section.page_height,
            section.left_margin,
            section.right_margin,
            section.top_margin,
            section.bottom_margin,
            section.header_distance,
            section.footer_distance,
        ]
    )


def has_page_field(paragraph) -> bool:
    values = paragraph._p.xpath(".//w:instrText/text()")
    return any("PAGE" in str(value).upper() for value in values)


def image_package_errors(doc: Document, docx_path: Path) -> list[str]:
    errors: list[str] = []
    referenced_rel_ids = set(doc._element.body.xpath(".//a:blip/@r:embed"))
    image_rels = {
        rel_id: rel
        for rel_id, rel in doc.part.rels.items()
        if rel.reltype == RT.IMAGE
    }
    orphan_rels = sorted(set(image_rels) - referenced_rel_ids)
    missing_rels = sorted(referenced_rel_ids - set(image_rels))
    if orphan_rels:
        errors.append(f"DOCX has residual image relationships: {', '.join(orphan_rels)}")
    if missing_rels:
        errors.append(f"DOCX image references are unresolved: {', '.join(missing_rels)}")
    targets = {"word/" + str(rel.target_ref).replace("\\", "/") for rel in image_rels.values()}
    with zipfile.ZipFile(docx_path) as archive:
        media = {name for name in archive.namelist() if name.startswith("word/media/") and not name.endswith("/")}
    if media != targets:
        extra = sorted(media - targets)
        missing = sorted(targets - media)
        if extra:
            errors.append(f"DOCX package has residual media: {', '.join(extra)}")
        if missing:
            errors.append(f"DOCX package is missing referenced media: {', '.join(missing)}")
    return errors


def pagination_control_errors(doc: Document) -> list[str]:
    found = {tag: [] for tag in PAGINATION_CONTROL_TAGS}
    for index, paragraph in enumerate(doc.paragraphs, start=1):
        ppr = paragraph._p.pPr
        if ppr is None:
            continue
        for tag in PAGINATION_CONTROL_TAGS:
            if ppr.find(qn(f"w:{tag}")) is not None:
                found[tag].append(index)
    details = [f"{tag}={len(indices)}" for tag, indices in found.items() if indices]
    if not details:
        return []
    return [
        "DOCX body contains black-square pagination controls: "
        + ", ".join(details)
    ]


def validate_docx(root: Path, docx_path: Path) -> list[str]:
    preflight, _warnings = validate_workspace(root)
    if preflight:
        return ["review package preflight failed", *preflight]
    if not docx_path.exists():
        return [f"DOCX not found: {docx_path}"]

    manifest = json.loads((root / "review_manifest.json").read_text(encoding="utf-8"))
    schema_version = str(manifest.get("schema_version", ""))
    report_type = str(manifest.get("report_type", ""))
    project_name = str(manifest.get("project_name", "")).strip()
    title, filename = report_identity(report_type, project_name)
    sections = SINGLE_SECTIONS if report_type == "single" else SITE_SECTIONS
    template_path = SINGLE_TEMPLATE if report_type == "single" else SITE_TEMPLATE
    reportable_status = "ai_ready" if schema_version == "1.6" else "verified"
    issues = [row for row in rows(root / "issue_candidates.csv") if row.get("status", "").strip() == reportable_status]
    issues.sort(key=lambda row: int(row.get("display_order", "0") or 0))

    doc = Document(docx_path)
    text = document_text(doc)
    errors: list[str] = []
    if docx_path.name != filename:
        errors.append(f"DOCX filename must be exactly: {filename}")
    if not doc.paragraphs or doc.paragraphs[0].text != title:
        errors.append(f"DOCX first paragraph title must be exactly: {title}")
    if len(doc.paragraphs) < 2 or not re.fullmatch(r"\d{4}年\d{1,2}月\d{1,2}日", doc.paragraphs[1].text.strip()):
        errors.append("DOCX second paragraph must be a Chinese date")
    if doc.paragraphs:
        title_p = doc.paragraphs[0]
        if title_p.alignment != WD_ALIGN_PARAGRAPH.CENTER:
            errors.append("DOCX title must be centered")
        if not title_p.runs or title_p.runs[0].font.size is None or round(title_p.runs[0].font.size.pt, 1) != 18.0:
            errors.append("DOCX title must use 18 pt font")
        if not title_p.runs or title_p.runs[0].bold is not True:
            errors.append("DOCX title must be bold")
    if text.count(title) != 1:
        errors.append("DOCX title must appear exactly once")
    errors.extend(pagination_control_errors(doc))

    cursor = -1
    grouped = {key: [] for key, _heading in sections}
    for issue in issues:
        grouped[issue.get("report_section", "").strip()].append(issue)
    for key, heading in sections:
        expected = heading + ("无意见" if not grouped[key] else "")
        matches = [index for index, paragraph in enumerate(doc.paragraphs) if paragraph.text == expected]
        if len(matches) != 1:
            errors.append(f"DOCX section heading must appear exactly once: {expected}")
            continue
        position = matches[0]
        if position <= cursor:
            errors.append(f"DOCX section heading is out of order: {expected}")
        cursor = position
        paragraph = doc.paragraphs[position]
        if not paragraph.runs or paragraph.runs[0].bold is not True:
            errors.append(f"DOCX section heading must be bold: {expected}")
        if report_type == "single":
            num_ids = paragraph._p.xpath("./w:pPr/w:numPr/w:numId/@w:val")
            if num_ids != ["7"]:
                errors.append(f"single-building section must use formal template numbering: {expected}")

    if text.count("【审查意见】：") != len(issues):
        errors.append(f"DOCX opinion count {text.count('【审查意见】：')} does not match reportable issue count {len(issues)}")
    opinion_cursor = -1
    for number, issue in enumerate(issues, start=1):
        issue_id = issue.get("issue_id", "").strip() or f"item {number}"
        markers = [
            f"{number}、涉及图纸：{issue.get('drawing_refs', '').strip()}",
            f"【审查意见】：{issue.get('problem', '').strip()}",
            f"【法规条文】：{regulation_text(issue, schema_version)}",
            f"【意见类型】：{finish_sentence(issue.get('opinion_type', '').strip())}",
        ]
        for marker in markers:
            if marker not in text:
                errors.append(f"{issue_id}: DOCX missing content: {marker}")
        position = text.find(markers[1], opinion_cursor + 1)
        if position < 0:
            errors.append(f"{issue_id}: review opinion is missing or out of order")
        else:
            opinion_cursor = position

    for issue in rows(root / "issue_candidates.csv"):
        if issue.get("status", "").strip() == reportable_status:
            continue
        problem = issue.get("problem", "").strip()
        if problem and f"【审查意见】：{problem}" in text:
            errors.append(f"non-reportable issue appears in DOCX: {issue.get('issue_id', '[missing issue_id]')}")

    expected_pictures = sum(len(resolve_paths(root, issue.get("screenshot_path", ""))) for issue in issues)
    if len(doc.inline_shapes) != expected_pictures:
        errors.append(f"DOCX picture count {len(doc.inline_shapes)} does not match expected {expected_pictures}")
    for marker in FORBIDDEN_PROCESS_TEXT:
        if marker in text:
            errors.append(f"DOCX contains forbidden process text: {marker}")
    for marker in ["图纸证据", "证据截图", "截图说明"]:
        if marker in text:
            errors.append(f"DOCX evidence image must not have a caption: {marker}")
    if text.count(FOOTER_NOTE) != 1:
        errors.append("DOCX must contain the exact formal 12-type opinion note once")
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("【法规条文】：") and ".pdf" in paragraph.text.casefold():
            errors.append("DOCX regulation display leaks a .pdf filename")

    template = Document(template_path)
    if not doc.sections or section_metrics(doc.sections[0]) != section_metrics(template.sections[0]):
        errors.append("DOCX page settings do not match the formal template")
    for index, section in enumerate(doc.sections, start=1):
        if not any(has_page_field(paragraph) for paragraph in section.footer.paragraphs):
            errors.append(f"DOCX section {index} footer is missing a PAGE field")
    errors.extend(image_package_errors(doc, docx_path))
    return errors


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("project_workspace", type=Path)
    parser.add_argument("docx", type=Path)
    args = parser.parse_args()
    errors = validate_docx(args.project_workspace.resolve(), args.docx.resolve())
    if errors:
        for error in errors:
            print(f"FAIL: {error}")
        return 1
    print("PASS: DOCX matches the formal AI initial review contract")
    return 0
