#!/usr/bin/env python3
"""Generate a formal-template AI initial review DOCX from a validated workspace."""

from __future__ import annotations

import argparse
import json
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt

from validate_review_package import image_size, resolve_paths, rows, validate_workspace

SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_DIR = SCRIPT_DIR.parent
SINGLE_TEMPLATE = SKILL_DIR / "范本_建筑单体施工图审查意见.docx"
SITE_TEMPLATE = SKILL_DIR / "范本_建筑总图施工图审查意见.docx"

SINGLE_SECTIONS = [
    ("设计说明", "设计说明："),
    ("平面图", "平面图："),
    ("立面剖面图", "立面、剖面图："),
    ("大样图", "大样图："),
]
SITE_SECTIONS = [
    ("总图设计说明", "总图（建筑）施工图设计说明："),
    ("总图设计图纸", "总图设计图纸："),
]
FOOTER_NOTE = (
    "注：【意见类型】包含消防安全强制性条文、一般性条文、政策规定、设计深度与必须修改（消防安全）"
    "组合4种类型；其它强制性条文、一般性条文、政策规定、设计深度与必须修改（其它）、建议修改（其它）"
    "的两两组合共8种类型，共12种意见类型。"
)


def set_font(run, name: str, size: float, bold: bool = False) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), name)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:eastAsia"), "zh-CN")


def configure_paragraph(
    paragraph,
    *,
    alignment=WD_ALIGN_PARAGRAPH.LEFT,
    before: float = 0,
    after: float = 4,
    line_spacing: float = 1.15,
    keep_next: bool = False,
    keep_together: bool = True,
) -> None:
    paragraph.alignment = alignment
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing
    fmt.keep_with_next = keep_next
    fmt.keep_together = keep_together
    fmt.widow_control = True


def set_numbering(paragraph, num_id: int = 7) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    numpr = ppr.find(qn("w:numPr"))
    if numpr is None:
        numpr = OxmlElement("w:numPr")
        ppr.insert(0, numpr)
    for child in list(numpr):
        numpr.remove(child)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), str(num_id))
    numpr.extend([ilvl, numid])


def clear_reference_body(doc: Document) -> None:
    body = doc._element.body
    sectpr = body.sectPr
    for child in list(body):
        if child is not sectpr:
            body.remove(child)
    for rel_id, rel in list(doc.part.rels.items()):
        if rel.reltype == RT.IMAGE:
            doc.part.drop_rel(rel_id)


def add_plain_paragraph(
    doc: Document,
    text: str,
    *,
    font: str = "宋体",
    size: float = 12,
    bold: bool = False,
    alignment=WD_ALIGN_PARAGRAPH.LEFT,
    before: float = 0,
    after: float = 4,
    keep_next: bool = False,
):
    paragraph = doc.add_paragraph()
    configure_paragraph(
        paragraph,
        alignment=alignment,
        before=before,
        after=after,
        keep_next=keep_next,
    )
    run = paragraph.add_run(text)
    set_font(run, font, size, bold)
    return paragraph


def add_drawing_reference(doc: Document, number: int, text: str) -> None:
    paragraph = doc.add_paragraph()
    configure_paragraph(paragraph, after=3, keep_next=True)
    prefix = paragraph.add_run(f"{number}、涉及图纸：")
    set_font(prefix, "宋体", 12, False)
    cursor = 0
    while cursor < len(text):
        start = text.find("《", cursor)
        if start < 0:
            run = paragraph.add_run(text[cursor:])
            set_font(run, "宋体", 12, False)
            break
        if start > cursor:
            run = paragraph.add_run(text[cursor:start])
            set_font(run, "宋体", 12, False)
        end = text.find("》", start)
        if end < 0:
            run = paragraph.add_run(text[start:])
            set_font(run, "宋体", 12, True)
            break
        run = paragraph.add_run(text[start : end + 1])
        set_font(run, "宋体", 12, True)
        cursor = end + 1


def add_labeled_paragraph(
    doc: Document,
    label: str,
    text: str,
    *,
    body_bold: bool = False,
    after: float = 4,
    keep_next: bool = False,
) -> None:
    paragraph = doc.add_paragraph()
    configure_paragraph(paragraph, after=after, keep_next=keep_next)
    label_run = paragraph.add_run(label)
    set_font(label_run, "宋体", 12, True)
    body_run = paragraph.add_run(text)
    set_font(body_run, "宋体", 12, body_bold)


def fit_picture_width(path: Path, max_width: float = 6.05, max_height: float = 4.6) -> Inches:
    size = image_size(path)
    if not size:
        return Inches(max_width)
    width, height = size
    if width <= 0 or height <= 0:
        return Inches(max_width)
    return Inches(min(max_width, max_height * width / height))


def add_evidence(doc: Document, path: Path, *, keep_next: bool) -> None:
    paragraph = doc.add_paragraph()
    configure_paragraph(
        paragraph,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        before=1,
        after=5,
        line_spacing=1.0,
        keep_next=keep_next,
    )
    paragraph.add_run().add_picture(str(path), width=fit_picture_width(path))


def add_page_number(paragraph) -> None:
    configure_paragraph(
        paragraph,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        after=0,
        line_spacing=1.0,
        keep_together=False,
    )
    run = paragraph.add_run()
    set_font(run, "宋体", 9, False)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def configure_footer_and_fields(doc: Document) -> None:
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")
    for section in doc.sections:
        section.footer_distance = Inches(0.5)
        footer = section.footer
        for paragraph in list(footer.paragraphs)[1:]:
            footer._element.remove(paragraph._element)
        paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        for child in list(paragraph._p):
            if child.tag != qn("w:pPr"):
                paragraph._p.remove(child)
        add_page_number(paragraph)


def add_section_heading(doc: Document, text: str, *, numbered: bool, empty: bool) -> None:
    paragraph = doc.add_paragraph()
    if numbered:
        set_numbering(paragraph)
    configure_paragraph(paragraph, before=6, after=6, keep_next=not empty)
    run = paragraph.add_run(text + ("无意见" if empty else ""))
    set_font(run, "黑体", 12, True)


def finish_sentence(text: str) -> str:
    text = text.rstrip()
    return text if text.endswith(("。", "！", "？", ".", "!", "?")) else text + "。"


def regulation_text(issue: dict[str, str], schema_version: str) -> str:
    if schema_version in {"1.2", "1.3", "1.4", "1.5", "1.6"} and issue.get("citation_mode", "").strip().lower() == "none":
        return "无。"
    source = issue.get("standard_display_name", "").strip()
    if not source:
        source_value = issue.get("standard_source", "").strip()
        source = Path(source_value).name if source_value else ""
    article = issue.get("standard_article", "").strip()
    requirement = issue.get("standard_requirement", "").strip()
    text = f"{source}{article}"
    if requirement:
        text += f"：{requirement}"
    return finish_sentence(text)


def safe_filename_component(value: str) -> str:
    value = re.sub(r'[\\/:*?"<>|]+', "-", value.strip())
    return value or "未命名项目"


def chinese_date(value: date) -> str:
    return f"{value.year}年{value.month}月{value.day}日"


def load_manifest(root: Path) -> tuple[str, dict]:
    manifest = json.loads((root / "review_manifest.json").read_text(encoding="utf-8"))
    return str(manifest.get("schema_version", "")), manifest


def report_identity(report_type: str, project_name: str) -> tuple[str, str]:
    project = safe_filename_component(project_name)
    suffix = "建筑施工图审查意见" if report_type == "single" else "建筑总图施工图审查意见"
    title = f"【AI初审】{project}{suffix}"
    return title, f"{title}.docx"


def build_report(args: argparse.Namespace) -> Path:
    root = args.project_workspace.resolve()
    errors, _warnings = validate_workspace(root)
    if errors:
        raise RuntimeError("review package validation failed:\n- " + "\n- ".join(errors))
    schema_version, manifest = load_manifest(root)
    report_type = str(manifest.get("report_type", ""))
    if args.report_type and args.report_type != report_type:
        raise RuntimeError(f"--report-type {args.report_type} conflicts with manifest report_type {report_type}")
    project_name = args.project_name or str(manifest.get("project_name", "")).strip()
    if not project_name:
        raise RuntimeError("project_name is required in review_manifest.json or --project-name")
    if args.project_name and manifest.get("project_name") and args.project_name != manifest.get("project_name"):
        raise RuntimeError("--project-name conflicts with review_manifest project_name")
    if report_type not in {"single", "site"}:
        raise RuntimeError(f"unsupported report_type: {report_type}")

    template = SINGLE_TEMPLATE if report_type == "single" else SITE_TEMPLATE
    sections = SINGLE_SECTIONS if report_type == "single" else SITE_SECTIONS
    title, filename = report_identity(report_type, project_name)
    if args.report_title and args.report_title != title:
        raise RuntimeError(f"--report-title must be exactly: {title}")
    output = (args.output or root / "output" / filename).resolve()
    reportable_status = "ai_ready" if schema_version == "1.6" else "verified"
    issues = [row for row in rows(root / "issue_candidates.csv") if row.get("status", "").strip() == reportable_status]
    issues.sort(key=lambda row: int(row.get("display_order", "0") or 0))
    grouped = {key: [] for key, _heading in sections}
    for issue in issues:
        grouped[issue.get("report_section", "").strip()].append(issue)

    doc = Document(template)
    clear_reference_body(doc)
    properties = doc.core_properties
    properties.title = title
    properties.subject = "建筑施工图审查意见"
    properties.author = ""
    properties.last_modified_by = ""
    properties.comments = ""

    add_plain_paragraph(
        doc,
        title,
        size=18,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        before=2,
        after=10,
        keep_next=True,
    )
    add_plain_paragraph(
        doc,
        args.report_date or chinese_date(date.today()),
        size=12,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        after=14,
        keep_next=True,
    )

    number = 1
    for section_key, heading in sections:
        section_issues = grouped[section_key]
        add_section_heading(doc, heading, numbered=report_type == "single", empty=not section_issues)
        for issue in section_issues:
            add_drawing_reference(doc, number, issue.get("drawing_refs", "").strip())
            add_labeled_paragraph(
                doc,
                "【审查意见】：",
                issue.get("problem", "").strip(),
                body_bold=True,
                after=5,
            )
            screenshots = resolve_paths(root, issue.get("screenshot_path", ""))
            for index, screenshot in enumerate(screenshots):
                add_evidence(doc, screenshot, keep_next=index == len(screenshots) - 1)
            add_labeled_paragraph(
                doc,
                "【法规条文】：",
                regulation_text(issue, schema_version),
                after=3,
                keep_next=True,
            )
            add_labeled_paragraph(
                doc,
                "【意见类型】：",
                finish_sentence(issue.get("opinion_type", "").strip()),
                after=8,
            )
            number += 1

    add_plain_paragraph(doc, FOOTER_NOTE, size=11.5, bold=True, before=8, after=0)
    configure_footer_and_fields(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    return output


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("project_workspace", type=Path)
    parser.add_argument("--report-type", choices=["single", "site"])
    parser.add_argument("--project-name")
    parser.add_argument("--report-title")
    parser.add_argument("--report-date", default="")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    try:
        print(build_report(args))
    except RuntimeError as error:
        print(f"FAIL: {error}")
        return 1
    return 0
