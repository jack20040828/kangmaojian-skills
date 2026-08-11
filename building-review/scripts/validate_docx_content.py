#!/usr/bin/env python3
"""Validate that a generated review DOCX faithfully contains verified issues."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document

from validate_review_package import resolve_paths, rows, validate_workspace

SECTION_HEADINGS_V11 = {
    "设计说明": "一、设计说明：",
    "平面图": "二、平面图：",
    "立面剖面图": "三、立面、剖面图：",
    "大样图": "四、大样图：",
    "总图设计说明": "总图（建筑）施工图设计说明：",
    "总图设计图纸": "总图设计图纸：",
}
SECTION_HEADINGS_V12 = dict(SECTION_HEADINGS_V11, 设计说明="设计说明：")


def document_text(doc: Document) -> str:
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def finish_sentence(text: str) -> str:
    text = text.rstrip()
    return text if text.endswith(("。", "！", "？", ".", "!", "?")) else text + "。"


def validate_docx(root: Path, docx_path: Path) -> list[str]:
    errors, _warnings = validate_workspace(root)
    if errors:
        return ["review package preflight failed"] + errors
    if not docx_path.exists():
        return [f"DOCX not found: {docx_path}"]
    issues = rows(root / "issue_candidates.csv")
    verified = [row for row in issues if row.get("status", "").strip() == "verified"]
    manifest_path = root / "review_manifest.json"
    schema_version = ""
    if manifest_path.exists():
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        schema_version = str(manifest.get("schema_version", ""))
    versioned = schema_version in {"1.1", "1.2"}
    if versioned:
        verified.sort(key=lambda row: int(row["display_order"]))

    doc = Document(docx_path)
    text = document_text(doc)
    errors = []
    opinion_marker_count = text.count("【审查意见】：")
    if opinion_marker_count != len(verified):
        errors.append(f"DOCX opinion count {opinion_marker_count} does not match verified issue count {len(verified)}")
    if not verified:
        if "经审查，本次未形成正式审查意见。" not in text:
            errors.append("zero-opinion DOCX is missing the required statement")
    else:
        cursor = -1
        for number, issue in enumerate(verified, start=1):
            issue_id = issue.get("issue_id", "") or f"item {number}"
            if schema_version == "1.2" and issue.get("citation_mode", "").strip().lower() == "none":
                citation_markers = ["【法规条文】：无。"]
            elif schema_version == "1.2":
                citation_markers = [
                    issue.get("standard_display_name", "").strip(),
                    issue.get("standard_article", "").strip(),
                    issue.get("standard_requirement", "").strip(),
                ]
            else:
                citation_markers = [
                    Path(issue.get("standard_source", "").strip()).name,
                    issue.get("standard_article", "").strip(),
                    issue.get("standard_requirement", "").strip(),
                ]
            markers = [
                finish_sentence(f"{number}、涉及图纸：{issue.get('drawing_refs', '').strip()}"),
                f"【审查意见】：{issue.get('problem', '').strip()}",
                *citation_markers,
                finish_sentence(f"【意见类型】：{issue.get('opinion_type', '').strip()}"),
            ]
            for marker in markers:
                if marker and marker not in text:
                    errors.append(f"{issue_id}: DOCX missing content: {marker}")
            position = text.find(markers[1], cursor + 1)
            if position < 0:
                errors.append(f"{issue_id}: review opinion is missing or out of order")
            else:
                cursor = position
            if versioned:
                heading_map = SECTION_HEADINGS_V12 if schema_version == "1.2" else SECTION_HEADINGS_V11
                heading = heading_map.get(issue.get("report_section", "").strip(), "")
                if heading and text.find(heading) > position:
                    errors.append(f"{issue_id}: appears before its declared report section")

    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("【法规条文】：") and ".pdf" in paragraph.text.casefold():
            errors.append("DOCX regulation display leaks a .pdf filename")

    for issue in issues:
        if issue.get("status", "").strip() == "verified":
            continue
        problem = issue.get("problem", "").strip()
        if problem and f"【审查意见】：{problem}" in text:
            errors.append(f"non-verified issue appears in DOCX: {issue.get('issue_id', '[missing issue_id]')}")

    expected_pictures = sum(len(resolve_paths(root, issue.get("screenshot_path", ""))) for issue in verified)
    if len(doc.inline_shapes) != expected_pictures:
        errors.append(f"DOCX picture count {len(doc.inline_shapes)} does not match expected {expected_pictures}")
    return errors


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("project_workspace", type=Path)
    parser.add_argument("docx", type=Path)
    args = parser.parse_args()
    errors = validate_docx(args.project_workspace.resolve(), args.docx.resolve())
    if errors:
        for error in errors:
            print(f"FAIL: {error}")
        return 1
    print("PASS: DOCX content matches verified review issues")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
