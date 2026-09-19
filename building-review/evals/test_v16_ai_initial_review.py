#!/usr/bin/env python3
"""Exercise v1.6 autonomous gates and produce anonymous single/site DOCX fixtures."""

from __future__ import annotations

import csv
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.oxml.ns import qn

SKILL_DIR = Path(__file__).resolve().parent.parent
SCRIPTS = SKILL_DIR / "scripts"
sys.path.insert(0, str(SCRIPTS))

from create_review_workspace import PROFILE_FACT_KEYS  # noqa: E402
from formal_review_report import PAGINATION_CONTROL_TAGS  # noqa: E402
from validate_review_package import V16_GATE_FIELDS  # noqa: E402


def run(args: list[str], expected: int = 0, contains: str = "") -> subprocess.CompletedProcess[str]:
    env = dict(os.environ)
    env["PYTHONDONTWRITEBYTECODE"] = "1"
    env["PYTHONIOENCODING"] = "utf-8"
    result = subprocess.run(args, text=True, encoding="utf-8", errors="replace", capture_output=True, env=env)
    output = result.stdout + result.stderr
    if result.returncode != expected or (contains and contains not in output):
        raise AssertionError(
            f"command failed expectation ({expected}, {contains!r}): {' '.join(args)}\n"
            f"exit={result.returncode}\n{output}"
        )
    return result


def read_rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def write_rows(path: Path, records: list[dict[str, str]], extra_fields: list[str] | None = None) -> None:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        fields = list(csv.DictReader(handle).fieldnames or [])
    for field in extra_fields or []:
        if field not in fields:
            fields.append(field)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(records)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def make_evidence(path: Path, label: str, box: tuple[int, int, int, int]) -> None:
    image = Image.new("RGB", (900, 520), "white")
    draw = ImageDraw.Draw(image)
    for x in range(80, 850, 110):
        draw.line((x, 60, x, 455), fill="#777777", width=2)
    for y in range(80, 450, 90):
        draw.line((55, y, 855, y), fill="#aaaaaa", width=2)
    draw.rectangle(box, outline="#d71920", width=8)
    draw.text((70, 25), label, fill="black")
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path)


def pagination_control_counts(path: Path) -> dict[str, int]:
    doc = Document(path)
    counts = {tag: 0 for tag in PAGINATION_CONTROL_TAGS}
    for paragraph in doc.paragraphs:
        ppr = paragraph._p.pPr
        if ppr is None:
            continue
        for tag in PAGINATION_CONTROL_TAGS:
            if ppr.find(qn(f"w:{tag}")) is not None:
                counts[tag] += 1
    return counts


def assert_no_pagination_controls(path: Path) -> None:
    present = {
        tag: count
        for tag, count in pagination_control_counts(path).items()
        if count
    }
    if present:
        raise AssertionError(f"DOCX contains black-square pagination controls: {present}")


def negative_docx_pagination_gate(workspace: Path, source_docx: Path, root: Path) -> None:
    bad_dir = root / "bad-pagination-docx"
    bad_dir.mkdir(parents=True, exist_ok=True)
    bad_docx = bad_dir / source_docx.name
    doc = Document(source_docx)
    doc.paragraphs[0].paragraph_format.keep_together = True
    doc.save(bad_docx)
    run(
        [sys.executable, str(SCRIPTS / "validate_docx_content.py"), str(workspace), str(bad_docx)],
        1,
        "black-square pagination controls",
    )


def make_rule_assets(root: Path) -> tuple[Path, Path, Path]:
    standard = root / "knowledge" / "B_核心规范" / "匿名建筑技术标准.txt"
    standard.parent.mkdir(parents=True, exist_ok=True)
    standard.write_text("匿名建筑技术标准 第1.0.1条：匿名图示应满足示例技术要求。\n", encoding="utf-8")
    relative = "B_核心规范/匿名建筑技术标准.txt"
    catalog = {
        "schema_version": "1.1",
        "rules": [{
            "rule_id": "ANON-1.0.1",
            "status": "active",
            "rule_pack": "anonymous_v16",
            "specialty": "001建筑设计文件编制深度专项",
            "report_types": ["single", "site"],
            "review_families": ["其他"],
            "coverage_topic": "document_traceability",
            "authority_mode": "normative",
            "trigger": {"always": True},
            "applicability_conditions": [{
                "condition_id": "C1",
                "description": "匿名图纸属于本次审查范围",
                "expected": True,
                "outcome_when_false": "not_applicable",
            }],
            "required_facts": ["匿名图纸事实"],
            "comparison_method": "将匿名图纸事实与匿名条文逐项比较。",
            "calculation_required": False,
            "graphic_evidence_requirements": {
                "required": False,
                "claim_type": "none",
                "required_roles": [],
                "sensitive_to_ambiguity": False,
            },
            "independent_review_required": True,
            "basis": {
                "source_role": "B_核心规范",
                "relative_path": relative,
                "sha256": sha256(standard),
                "article": "第1.0.1条",
                "requirement": "匿名图示应满足示例技术要求。",
            },
        }],
    }
    catalog_path = root / "anonymous-review-rules.json"
    catalog_path.write_text(json.dumps(catalog, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    index = {
        "schema_version": "1.1",
        "generated_at": "2026-08-24T09:00:00",
        "corpus_sha256": sha256(standard),
        "total_files": 1,
        "specialty_count": 1,
        "files": [{
            "id": "K0001",
            "specialty": "001建筑设计文件编制深度专项",
            "role": "B_核心规范",
            "kind": "txt",
            "name": standard.name,
            "relative_path": relative,
            "absolute_path": str(standard),
            "size_bytes": standard.stat().st_size,
            "sha256": sha256(standard),
            "text_available": True,
            "extraction_status": "text",
            "text_sample": "匿名建筑技术标准",
        }],
    }
    index_path = root / "anonymous-knowledge-index.json"
    index_path.write_text(json.dumps(index, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return standard, catalog_path, index_path


def complete_profile(path: Path) -> None:
    profile = json.loads(path.read_text(encoding="utf-8"))
    boolean_fields = {
        "industrial_building", "overnight_stay", "sprinkler", "basement", "elevator",
        "accessible_requirement", "roof_accessible", "wet_rooms", "parking",
        "photovoltaic_or_solar", "food_service", "dormitory_or_hotel", "school",
        "office", "children_activity",
    }
    values = {
        "location_province": "湖南省",
        "location_city": "匿名市",
        "building_use": "办公楼",
        "gross_floor_area_m2": 800,
        "building_height_m": 9.0,
        "floors_above": 3,
        "floors_below": 0,
        "fire_hazard_class": "不适用",
        "fire_resistance_rating": "二级",
        "occupant_load": 30,
    }
    if set(profile["facts"]) != set(PROFILE_FACT_KEYS):
        raise AssertionError("v1.6 profile fact keys changed unexpectedly")
    for key, record in profile["facts"].items():
        record["status"] = "ai_completed"
        record["value"] = False if key in boolean_fields else values.get(key, "已由匿名图纸确认")
        record["fact_ids"] = ["F001"]
        record["notes"] = "AI由匿名图纸事实关闭。"
    profile["ai_review_completed"] = True
    profile["completed_at"] = "2026-08-24T10:00:00"
    profile["route_completion"] = {
        "status": "AI初审完成",
        "completed_at": "2026-08-24T09:50:00",
        "notes": "AI已依据匿名项目事实关闭专项路由。",
    }
    profile["discovery_tracks"] = {
        "technical_compliance": {"status": "completed", "notes": "AI已完成技术条文检查。"},
        "design_depth": {"status": "completed", "notes": "AI已完成设计深度与跨图检查。"},
        "optimization": {"status": "completed", "notes": "AI已完成有依据的优化扫描。"},
    }
    path.write_text(json.dumps(profile, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def configure_manifest(workspace: Path, catalog: Path) -> None:
    manifest_path = workspace / "review_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["rule_catalog_snapshot"] = {"path": str(catalog), "sha256": sha256(catalog)}
    manifest["knowledge_snapshot"] = {
        "required_layers": ["A", "B", "C", "D"],
        "layers": {
            layer: {
                "status": "used",
                "sources": [f"{layer}_匿名资料"],
                "notes": {
                    "A": "用于展开匿名审查主题。",
                    "B": "用于正式条文引用。",
                    "C": "用于确认匿名条文适用性。",
                    "D": "仅用于辅助发现和差异比对。",
                }[layer],
            }
            for layer in ["A", "B", "C", "D"]
        },
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def build_workspace(root: Path, report_type: str, catalog: Path, index: Path) -> tuple[Path, Path]:
    project_name = "匿名单体项目" if report_type == "single" else "匿名项目"
    created = run([
        sys.executable,
        str(SCRIPTS / "create_review_workspace.py"),
        project_name,
        "--root",
        str(root / "reviews"),
        "--schema-version",
        "1.6",
        "--report-type",
        report_type,
    ])
    workspace = Path(created.stdout.strip().splitlines()[-1])
    if (workspace / "independent_review_log.csv").exists():
        raise AssertionError("v1.6 workspace must not create independent_review_log.csv")
    for ledger in ["check_matrix.csv", "validation_log.csv", "graphic_evidence_chain.csv", "applicability_decisions.csv"]:
        with (workspace / ledger).open(encoding="utf-8-sig", newline="") as handle:
            fields = list(csv.DictReader(handle).fieldnames or [])
        if any(field in fields for field in ["reviewer_name", "reviewer_confirmation", "reviewed_at", "reviewer_gate", "independent_review_id"]):
            raise AssertionError(f"v1.6 ledger leaks a human-review field: {ledger}")

    drawing = workspace / "source" / "anonymous-drawing.txt"
    drawing.write_text("anonymous drawing source\n", encoding="utf-8")
    drawing_name = "匿名平面图" if report_type == "single" else "匿名总平面图"
    inventory = {
        "source_file": "source/anonymous-drawing.txt",
        "page": "1",
        "sheet_no": "1",
        "drawing_no": "",
        "drawing_name": drawing_name,
        "discipline": "建筑",
        "content_type": "匿名图纸",
        "scale": "1:100",
        "title_block_status": "已脱敏",
        "review_family": "其他",
        "review_status": "needs_review",
        "review_check_ids": "",
        "review_notes": "AI原子检查尚未关闭。",
        "notes": "",
    }
    fact = {
        "fact_id": "F001",
        "source_file": "source/anonymous-drawing.txt",
        "page": "1",
        "drawing_no": "",
        "drawing_name": drawing_name,
        "location": "匿名图纸中部",
        "fact_type": "匿名图纸事实",
        "raw_text_or_measure": "匿名图示与示例技术要求不一致。" if report_type == "single" else "匿名图示符合示例技术要求。",
        "value": "",
        "unit": "",
        "confidence": "高",
        "needs_verification": "否",
        "notes": "",
    }
    write_rows(workspace / "drawing_inventory.csv", [inventory])
    write_rows(workspace / "fact_ledger.csv", [fact])
    complete_profile(workspace / "project_profile.json")
    configure_manifest(workspace, catalog)
    run([sys.executable, str(SCRIPTS / "generate_project_checklist.py"), str(workspace), "--catalog", str(catalog)])

    checks = read_rows(workspace / "check_matrix.csv")
    if len(checks) != 1:
        raise AssertionError("anonymous v1.6 catalog must expand exactly one check")
    check = checks[0]
    check.update(
        decision_state="resolved",
        applicability="适用",
        actual_fact=fact["raw_text_or_measure"],
        fact_ids="F001",
        drawing_refs=f"PDF第1页《{drawing_name}》",
        comparison_record="AI已把匿名图纸事实与匿名条文逐项比较。",
        calculation_record="",
        conclusion="不符合" if report_type == "single" else "符合",
        forms_issue="yes" if report_type == "single" else "no",
        issue_id="AI-001" if report_type == "single" else "",
        not_forming_reason="" if report_type == "single" else "图纸事实与匿名条文比较后符合。",
        open_reason="",
        graphic_claim_type="none",
        graphic_gate_reason="本匿名规则仅依据明确文字事实，不需要图形解释链。",
        completion_gate="AI初审完成",
        notes="AI已关闭本项原子检查。",
    )
    write_rows(workspace / "check_matrix.csv", [check])
    inventory["review_status"] = "reviewed"
    inventory["review_check_ids"] = check["check_id"]
    inventory["review_notes"] = "AI已关闭全部触发原子检查。"
    write_rows(workspace / "drawing_inventory.csv", [inventory])

    applicability = read_rows(workspace / "applicability_decisions.csv")
    applicability[0].update(
        actual_value="true",
        result="met",
        fact_ids="F001",
        drawing_refs=f"PDF第1页《{drawing_name}》",
        status="resolved",
        completed_at="2026-08-24T10:10:00",
        notes="AI已由匿名图纸事实关闭适用条件。",
    )
    write_rows(workspace / "applicability_decisions.csv", applicability)
    write_rows(workspace / "graphic_evidence_chain.csv", [])

    if report_type == "single":
        shot1 = workspace / "screenshots" / "AI-001-1.png"
        shot2 = workspace / "screenshots" / "AI-001-2.png"
        make_evidence(shot1, "ANONYMOUS EVIDENCE 1", (310, 170, 515, 335))
        make_evidence(shot2, "ANONYMOUS EVIDENCE 2", (485, 255, 725, 420))
        issue = {
            "issue_id": "AI-001",
            "status": "ai_ready",
            "display_order": "1",
            "report_section": "设计说明",
            "specialty": "001建筑设计文件编制深度专项",
            "check_id": check["check_id"],
            "fact_ids": "F001",
            "drawing_refs": f"PDF第1页《{drawing_name}》",
            "problem": "如本条适用于本项目，匿名图示与示例技术要求不一致，应按匿名条文补充并统一表达。",
            "citation_mode": "cited",
            "standard_source": "B_核心规范/匿名建筑技术标准.txt",
            "standard_display_name": "《匿名建筑技术标准》",
            "standard_article": "第1.0.1条",
            "standard_requirement": "匿名图示应满足示例技术要求。",
            "citation_none_reason": "",
            "judgment": "匿名图纸事实与条文要求直接比较后不一致。",
            "case_refs": "D_匿名案例仅作差异比对",
            "needs_screenshot": "yes",
            "screenshot_path": "screenshots/AI-001-1.png;screenshots/AI-001-2.png",
            "screenshot_location": "匿名图纸中部及右下部",
            "screenshot_strategy": "multiple",
            "screenshot_reason": "两处图示共同证明同一根因。",
            "screenshot_count": "2",
            "evidence_point": "匿名图示与条文要求的直接差异",
            "red_box_target": "两处匿名差异位置",
            "context_required": "保留周边网格和定位线",
            "screenshot_quality": "raster_high",
            "opinion_type": "其它强制性条文，必须修改（其它）",
            "validation_status": "AI初审完成",
            "notes": "AI已合并同一根因并确认该意见可交付。",
        }
        write_rows(workspace / "issue_candidates.csv", [issue])
        gate = {field: "" for field in read_header(workspace / "validation_log.csv")}
        for field in V16_GATE_FIELDS:
            gate[field] = "通过"
        gate.update(
            issue_id="AI-001",
            fact_check="通过",
            standard_check="通过",
            specialty_check="通过",
            case_check="通过",
            screenshot_check="通过",
            layout_check="通过",
            opinion_type_check="通过",
            gate_origin="agent",
            stage_completion="AI初审完成",
            completed_at="2026-08-24T10:20:00",
            result="通过",
            notes="AI已完成全部意见验证。",
        )
        write_rows(workspace / "validation_log.csv", [gate])
    else:
        write_rows(workspace / "issue_candidates.csv", [])
        write_rows(workspace / "validation_log.csv", [])

    run([
        sys.executable,
        str(SCRIPTS / "snapshot_review_integrity.py"),
        str(workspace),
        "--index",
        str(index),
        "--rule-catalog",
        str(catalog),
    ])
    run([sys.executable, str(SCRIPTS / "audit_review_completeness.py"), str(workspace)])
    run([sys.executable, str(SCRIPTS / "validate_review_package.py"), str(workspace)], contains="PASS:")
    report = run([sys.executable, str(SCRIPTS / "generate_review_report.py"), str(workspace)])
    docx = Path(report.stdout.strip().splitlines()[-1])
    run([sys.executable, str(SCRIPTS / "validate_docx_content.py"), str(workspace), str(docx)], contains="PASS:")
    return workspace, docx


def read_header(path: Path) -> list[str]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle).fieldnames or [])


def negative_gates(workspace: Path) -> None:
    validate = [sys.executable, str(SCRIPTS / "validate_review_package.py"), str(workspace)]
    issues = read_rows(workspace / "issue_candidates.csv")
    issues[0]["status"] = "verified"
    write_rows(workspace / "issue_candidates.csv", issues)
    run(validate, 1, "v1.6 uses ai_ready instead of verified")
    issues[0]["status"] = "ai_ready"
    write_rows(workspace / "issue_candidates.csv", issues)

    gates = read_rows(workspace / "validation_log.csv")
    gates[0]["reviewer_name"] = "虚构人工审查人"
    write_rows(workspace / "validation_log.csv", gates, extra_fields=["reviewer_name"])
    run(validate, 1, "v1.6 must not record reviewer_name")
    gates[0].pop("reviewer_name", None)
    write_rows(workspace / "validation_log.csv", gates)

    independent = workspace / "independent_review_log.csv"
    independent.write_text("review_id\n", encoding="utf-8-sig")
    run(validate, 1, "v1.6 must not create independent_review_log.csv")
    independent.unlink()

    gates = read_rows(workspace / "validation_log.csv")
    gates[0]["gate_origin"] = "manual"
    write_rows(workspace / "validation_log.csv", gates)
    run(validate, 1, "v1.6 gate_origin must be agent")
    gates[0]["gate_origin"] = "agent"
    write_rows(workspace / "validation_log.csv", gates)

    run([sys.executable, str(SCRIPTS / "audit_review_completeness.py"), str(workspace)])
    run(validate, 0, "PASS:")


def main() -> int:
    artifact_root_value = os.environ.get("BUILDING_REVIEW_V16_ARTIFACT_DIR", "").strip()
    with tempfile.TemporaryDirectory(prefix="building-review-v16-") as temp_value:
        root = Path(temp_value)
        _standard, catalog, index = make_rule_assets(root)
        single_workspace, single_docx = build_workspace(root / "single", "single", catalog, index)
        site_workspace, site_docx = build_workspace(root / "site", "site", catalog, index)
        assert_no_pagination_controls(single_docx)
        assert_no_pagination_controls(site_docx)
        negative_docx_pagination_gate(single_workspace, single_docx, root)
        negative_gates(single_workspace)

        if artifact_root_value:
            artifact_root = Path(artifact_root_value).resolve()
            artifact_root.mkdir(parents=True, exist_ok=True)
            for docx in [single_docx, site_docx]:
                shutil.copy2(docx, artifact_root / docx.name)
            (artifact_root / "fixture-summary.json").write_text(
                json.dumps(
                    {
                        "schema_version": "1.6",
                        "single": single_docx.name,
                        "site": site_docx.name,
                        "single_workspace": str(single_workspace),
                        "site_workspace": str(site_workspace),
                    },
                    ensure_ascii=False,
                    indent=2,
                ) + "\n",
                encoding="utf-8",
            )

    print("PASS: v1.6 autonomous single/site reports and negative human-gate tests completed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
