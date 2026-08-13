#!/usr/bin/env python3
"""Run anonymous end-to-end regression gates for building-review v1.1-v1.3."""

from __future__ import annotations

import base64
import csv
import json
import os
import subprocess
import sys
import tempfile
from contextlib import nullcontext
from datetime import date
from pathlib import Path

from docx import Document

SKILL_DIR = Path(__file__).resolve().parent.parent
SCRIPTS = SKILL_DIR / "scripts"
OPINION_TYPE = "一般性条文，必须修改（其它）"
DESIGN_DEPTH_TYPE = "设计深度，必须修改（其它）"
PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZQmcAAAAASUVORK5CYII="
)


def run(args: list[str], expected: int, contains: str = "") -> subprocess.CompletedProcess[str]:
    env = dict(os.environ)
    env["PYTHONDONTWRITEBYTECODE"] = "1"
    env["PYTHONIOENCODING"] = "utf-8"
    result = subprocess.run(args, text=True, encoding="utf-8", errors="replace", capture_output=True, env=env)
    output = result.stdout + result.stderr
    if result.returncode != expected or (contains and contains not in output):
        raise AssertionError(
            f"command failed expectation ({expected}, {contains!r}): {' '.join(args)}\n"
            f"exit={result.returncode}\n{output}"
        )
    return result


def headers(path: Path) -> list[str]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle).fieldnames or [])


def write_rows(path: Path, records: list[dict[str, str]]) -> None:
    columns = headers(path)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(records)


def issue(issue_id: str, check_id: str, order: str, standard: Path, status: str = "verified") -> dict[str, str]:
    return {
        "issue_id": issue_id,
        "status": status,
        "display_order": order,
        "report_section": "平面图",
        "specialty": "匿名专项",
        "check_id": check_id,
        "fact_ids": "F001",
        "drawing_refs": "A-01 一层平面图",
        "problem": f"匿名问题{issue_id}。",
        "citation_mode": "cited",
        "standard_source": str(standard),
        "standard_display_name": "《匿名规范》",
        "standard_article": "第1条",
        "standard_requirement": "匿名要求。",
        "citation_none_reason": "",
        "judgment": "图纸事实与要求不一致。",
        "case_refs": "无",
        "needs_screenshot": "no",
        "screenshot_path": "",
        "screenshot_location": "",
        "screenshot_strategy": "none",
        "screenshot_reason": "文字事实足以证明。",
        "screenshot_count": "0",
        "evidence_point": "",
        "red_box_target": "",
        "context_required": "",
        "screenshot_quality": "",
        "opinion_type": OPINION_TYPE,
        "validation_status": "通过",
        "notes": "已合并重复项并确认可交付。",
    }


def check(check_id: str, issue_id: str, forms_issue: str = "yes", conclusion: str = "不符合") -> dict[str, str]:
    return {
        "check_id": check_id,
        "specialty": "匿名专项",
        "review_item_id": "ANON-A-001",
        "coverage_topic": "function_layout",
        "evidence_class": "relationship",
        "source_review_item": "匿名审查要点",
        "applicability": "适用",
        "required_fact": "匿名要求事实",
        "actual_fact": "匿名实际事实",
        "fact_ids": "F001",
        "drawing_refs": "A-01",
        "standard_source": "匿名规范",
        "standard_article": "第1条",
        "standard_requirement": "匿名要求。",
        "conclusion": conclusion,
        "forms_issue": forms_issue,
        "issue_id": issue_id,
        "not_forming_reason": "" if forms_issue == "yes" else "图纸表达符合要求。",
        "notes": "",
    }


def gate(issue_id: str) -> dict[str, str]:
    return {
        "issue_id": issue_id,
        "fact_check": "通过",
        "standard_check": "通过",
        "specialty_check": "通过",
        "case_check": "通过",
        "screenshot_check": "通过",
        "professional_filter_check": "通过",
        "screenshot_strategy_check": "通过",
        "red_box_precision_check": "通过",
        "graphical_interpretation_check": "通过",
        "opinion_wording_check": "通过",
        "layout_check": "待检查",
        "opinion_type_check": "通过",
        "result": "通过",
        "notes": "",
    }


def set_schema_version(workspace: Path, version: str) -> None:
    path = workspace / "review_manifest.json"
    manifest = json.loads(path.read_text(encoding="utf-8"))
    manifest["schema_version"] = version
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def none_issue(issue_id: str, check_id: str, order: str) -> dict[str, str]:
    value = issue(issue_id, check_id, order, Path("unused"))
    value.update(
        citation_mode="none",
        standard_source="",
        standard_display_name="",
        standard_article="",
        standard_requirement="",
        citation_none_reason="图纸内部索引和表达完整性问题，不主张外部技术阈值。",
        opinion_type=DESIGN_DEPTH_TYPE,
    )
    return value


def main() -> int:
    metadata = json.loads((SKILL_DIR / "evals" / "regression-cases.json").read_text(encoding="utf-8"))
    if len(metadata) != 50:
        raise AssertionError("regression metadata must contain 50 cases")
    python = sys.executable
    keep_dir = os.environ.get("BUILDING_REVIEW_KEEP_TEMP", "")
    if keep_dir:
        Path(keep_dir).mkdir(parents=True, exist_ok=True)
        temp_context = nullcontext(keep_dir)
    else:
        temp_context = tempfile.TemporaryDirectory(prefix="building-review-v12-")
    with temp_context as temp_value:
        temp = Path(temp_value)
        knowledge = temp / "knowledge" / "001匿名专项" / "B_核心规范"
        knowledge.mkdir(parents=True)
        standard = knowledge / "匿名规范.txt"
        standard.write_text("第1条 匿名要求。\n", encoding="utf-8")
        index_dir = temp / "index"
        run([python, str(SCRIPTS / "build_knowledge_index.py"), "--knowledge-base", str(temp / "knowledge"), "--output-dir", str(index_dir)], 0)
        create = run([
            python,
            str(SCRIPTS / "create_review_workspace.py"),
            "匿名项目",
            "--root",
            str(temp / "reviews"),
            "--report-type",
            "single",
        ], 0)
        workspace = Path(create.stdout.strip().splitlines()[-1])
        set_schema_version(workspace, "1.1")
        drawing = workspace / "source" / "drawing.txt"
        drawing.write_text("anonymous drawing source\n", encoding="utf-8")
        base_inventory = {
            "source_file": "source/drawing.txt",
            "page": "1",
            "sheet_no": "1",
            "drawing_no": "A-01",
            "drawing_name": "一层平面图",
            "discipline": "建筑",
            "content_type": "平面图",
            "scale": "1:100",
            "title_block_status": "已核对",
            "review_family": "平面图",
            "review_status": "reviewed",
            "review_check_ids": "C001",
            "review_notes": "",
            "notes": "",
        }
        write_rows(workspace / "drawing_inventory.csv", [base_inventory])
        base_fact = {
            "fact_id": "F001",
            "source_file": "source/drawing.txt",
            "page": "1",
            "drawing_no": "A-01",
            "drawing_name": "一层平面图",
            "location": "匿名位置",
            "fact_type": "文字",
            "raw_text_or_measure": "匿名实际事实",
            "value": "",
            "unit": "",
            "confidence": "高",
            "needs_verification": "否",
            "notes": "",
        }
        base_check = check("C001", "I001")
        base_issue = issue("I001", "C001", "1", standard)
        base_gate = gate("I001")
        write_rows(workspace / "fact_ledger.csv", [base_fact])
        write_rows(workspace / "check_matrix.csv", [base_check])
        write_rows(workspace / "issue_candidates.csv", [base_issue])
        write_rows(workspace / "validation_log.csv", [base_gate])
        snapshot = [python, str(SCRIPTS / "snapshot_review_integrity.py"), str(workspace), "--index", str(index_dir / "knowledge-index.json")]
        validate = [python, str(SCRIPTS / "validate_review_package.py"), str(workspace)]
        run(snapshot, 0)
        run(validate, 0, "PASS:")

        write_rows(workspace / "fact_ledger.csv", [base_fact, dict(base_fact)])
        run(validate, 1, "duplicate fact_id")
        write_rows(workspace / "fact_ledger.csv", [base_fact])

        broken_issue = dict(base_issue, fact_ids="F999")
        write_rows(workspace / "issue_candidates.csv", [broken_issue])
        run(validate, 1, "fact_id not found")
        write_rows(workspace / "issue_candidates.csv", [base_issue])

        gap_issue = dict(base_issue, display_order="2")
        write_rows(workspace / "issue_candidates.csv", [gap_issue])
        run(validate, 1, "continuous from 1")
        write_rows(workspace / "issue_candidates.csv", [base_issue])

        second_issue = issue("I002", "C002", "1", standard)
        write_rows(workspace / "check_matrix.csv", [base_check, check("C002", "I002")])
        write_rows(workspace / "issue_candidates.csv", [base_issue, second_issue])
        write_rows(workspace / "validation_log.csv", [base_gate, gate("I002")])
        run(snapshot, 0)
        run(validate, 1, "duplicate display_order")
        write_rows(workspace / "check_matrix.csv", [base_check])
        write_rows(workspace / "issue_candidates.csv", [base_issue])
        write_rows(workspace / "validation_log.csv", [base_gate])
        run(snapshot, 0)

        drawing.write_text("changed source\n", encoding="utf-8")
        run(validate, 1, "source file changed")
        drawing.write_text("anonymous drawing source\n", encoding="utf-8")

        extra = workspace / "source" / "extra.txt"
        extra.write_text("extra\n", encoding="utf-8")
        run(validate, 1, "unregistered source")
        extra.unlink()

        standard.write_text("changed standard\n", encoding="utf-8")
        run(validate, 1, "used standard changed")
        standard.write_text("第1条 匿名要求。\n", encoding="utf-8")
        changed_source = dict(base_issue, standard_source="不存在的规范.txt")
        write_rows(workspace / "issue_candidates.csv", [changed_source])
        run(validate, 1, "not found in knowledge index")
        write_rows(workspace / "issue_candidates.csv", [base_issue])

        missing_shot = dict(
            base_issue,
            needs_screenshot="yes",
            screenshot_strategy="single",
            screenshot_path="screenshots/missing.png",
            screenshot_location="A-01",
            screenshot_count="1",
            evidence_point="匿名证据",
            red_box_target="匿名目标",
            context_required="保留图名",
            screenshot_quality="清晰",
        )
        write_rows(workspace / "issue_candidates.csv", [missing_shot])
        run(validate, 1, "does not exist")

        excluded = issue("I099", "C001", "", standard, status="needs_review")
        excluded["problem"] = "不得进入正式报告。"
        excluded["notes"] = "证据不足。"
        write_rows(workspace / "issue_candidates.csv", [base_issue, excluded])
        run(snapshot, 0)
        run(validate, 0)
        report = workspace / "output" / "report.docx"
        generate = [
            python,
            str(SCRIPTS / "generate_review_report.py"),
            str(workspace),
            "--project-name",
            "匿名项目",
            "--report-title",
            "建筑施工图审查意见",
            "--output",
            str(report),
        ]
        run(generate, 0)
        content = [python, str(SCRIPTS / "validate_docx_content.py"), str(workspace), str(report)]
        run(content, 0)

        second_issue = issue("I002", "C002", "2", standard)
        write_rows(workspace / "check_matrix.csv", [base_check, check("C002", "I002")])
        write_rows(workspace / "issue_candidates.csv", [base_issue, second_issue])
        write_rows(workspace / "validation_log.csv", [base_gate, gate("I002")])
        run(snapshot, 0)
        wrong_order = workspace / "output" / "wrong-order.docx"
        wrong = Document()
        wrong.add_paragraph("二、平面图：")
        for number, item in [(1, second_issue), (2, base_issue)]:
            wrong.add_paragraph(f"{number}、涉及图纸：{item['drawing_refs']}。")
            wrong.add_paragraph(f"【审查意见】：{item['problem']}")
            wrong.add_paragraph(f"【法规条文】：{standard.name}{item['standard_article']}：{item['standard_requirement']}。")
            wrong.add_paragraph(f"【意见类型】：{item['opinion_type']}。")
        wrong.save(wrong_order)
        run([python, str(SCRIPTS / "validate_docx_content.py"), str(workspace), str(wrong_order)], 1, "out of order")

        write_rows(workspace / "check_matrix.csv", [base_check])
        write_rows(workspace / "issue_candidates.csv", [base_issue, excluded])
        write_rows(workspace / "validation_log.csv", [base_gate])
        run(snapshot, 0)
        render_dir = workspace / "rendered"
        render_dir.mkdir()
        (render_dir / "page-1.png").write_bytes(PNG_1X1)
        qa = [python, str(SCRIPTS / "validate_report_qa.py"), str(workspace)]
        run(qa + ["--init", "--docx", str(report), "--render-dir", str(render_dir)], 0)
        run(qa, 1, "content_check is not")
        run(qa + ["--record-content-pass"], 0)
        run(qa + ["--record-visual-pass", "--checked-pages", "1"], 0)
        run(qa, 0, "PASS:")

        zero_check = check("C001", "", forms_issue="no", conclusion="符合")
        write_rows(workspace / "check_matrix.csv", [zero_check])
        write_rows(workspace / "issue_candidates.csv", [])
        write_rows(workspace / "validation_log.csv", [])
        run(snapshot, 0)
        run(validate, 0)
        zero_report = workspace / "output" / "zero.docx"
        run(generate[:-1] + [str(zero_report)], 0)
        run([python, str(SCRIPTS / "validate_docx_content.py"), str(workspace), str(zero_report)], 0)

        # v1.2: every identifiable sheet must close its review status.
        set_schema_version(workspace, "1.2")
        family_rows = []
        family_facts = [base_fact]
        family_checks = [base_check]
        for index, (drawing_no, drawing_name, family) in enumerate([
            ("A-01", "一层平面图", "平面图"),
            ("A-04", "屋面平面图", "屋面图"),
            ("A-07", "建筑剖面图", "剖面图"),
            ("A-08", "楼梯大样图", "楼梯大样"),
            ("A-09", "墙身大样图", "墙身大样"),
        ], start=1):
            check_id = "C001" if drawing_no == "A-01" else f"C1{index:02d}"
            family_rows.append(
                dict(
                    base_inventory,
                    drawing_no=drawing_no,
                    drawing_name=drawing_name,
                    review_family=family,
                    review_status="reviewed",
                    review_check_ids=check_id,
                )
            )
            if drawing_no != "A-01":
                fact_id = f"F1{index:02d}"
                family_facts.append(
                    dict(
                        base_fact,
                        fact_id=fact_id,
                        drawing_no=drawing_no,
                        drawing_name=drawing_name,
                        raw_text_or_measure="图纸标注控制尺寸1200mm。",
                    )
                )
                closure = check(check_id, "", forms_issue="no", conclusion="符合")
                closure.update(
                    fact_ids=fact_id,
                    drawing_refs=f"{drawing_no} {drawing_name}",
                    actual_fact="图纸标注控制尺寸1200mm。",
                    not_forming_reason="已核对图纸标注的1200mm控制尺寸，符合匿名要求。",
                )
                family_checks.append(closure)
        write_rows(workspace / "drawing_inventory.csv", family_rows)
        write_rows(workspace / "fact_ledger.csv", family_facts)
        write_rows(workspace / "check_matrix.csv", family_checks)
        write_rows(workspace / "issue_candidates.csv", [base_issue])
        write_rows(workspace / "validation_log.csv", [base_gate])
        run(snapshot, 0)
        run(validate, 0)

        missing_coverage = list(family_rows)
        missing_coverage[1] = dict(missing_coverage[1], review_status="")
        write_rows(workspace / "drawing_inventory.csv", missing_coverage)
        run(validate, 1, "invalid review_status")

        pending_coverage = list(family_rows)
        pending_coverage[2] = dict(pending_coverage[2], review_status="needs_review")
        write_rows(workspace / "drawing_inventory.csv", pending_coverage)
        run(validate, 1, "needs_review blocks report generation")
        write_rows(workspace / "drawing_inventory.csv", family_rows)

        generic_facts = list(family_facts)
        generic_facts[1] = dict(generic_facts[1], raw_text_or_measure="已提供屋面大样，覆盖完整。")
        generic_checks = list(family_checks)
        generic_checks[1] = dict(
            generic_checks[1],
            actual_fact="已提供屋面大样，覆盖完整。",
            not_forming_reason="图纸已提供，判定符合。",
        )
        write_rows(workspace / "fact_ledger.csv", generic_facts)
        write_rows(workspace / "check_matrix.csv", generic_checks)
        run(validate, 1, "generic drawing-presence statement")

        wrong_link_rows = list(family_rows)
        wrong_link_rows[1] = dict(wrong_link_rows[1], review_check_ids="C001")
        write_rows(workspace / "drawing_inventory.csv", wrong_link_rows)
        write_rows(workspace / "fact_ledger.csv", family_facts)
        write_rows(workspace / "check_matrix.csv", family_checks)
        run(validate, 1, "no substantive sheet-specific check")

        write_rows(workspace / "drawing_inventory.csv", [base_inventory])
        write_rows(workspace / "fact_ledger.csv", [base_fact])
        conditional_check = dict(base_check, applicability="需判断", conclusion="需核验")
        write_rows(workspace / "check_matrix.csv", [conditional_check])
        write_rows(workspace / "issue_candidates.csv", [base_issue])
        run(validate, 1, "unresolved check cannot support a verified issue")
        write_rows(workspace / "check_matrix.csv", [base_check])

        valid_none = none_issue("I001", "C001", "1")
        write_rows(workspace / "issue_candidates.csv", [valid_none])
        run(snapshot, 0)
        run(validate, 0)

        invalid_none = dict(valid_none, opinion_type=OPINION_TYPE)
        write_rows(workspace / "issue_candidates.csv", [invalid_none])
        run(validate, 1, "allowed only for design-depth opinion types")

        opinion_type_typo = dict(base_issue, opinion_type="设计深度文，必须修改（其它）")
        write_rows(workspace / "issue_candidates.csv", [opinion_type_typo])
        run(validate, 1, "invalid opinion_type")

        write_rows(workspace / "issue_candidates.csv", [base_issue])
        run(snapshot, 0)
        ambiguous_gate = dict(base_gate, graphical_interpretation_check="需核验")
        write_rows(workspace / "validation_log.csv", [ambiguous_gate])
        run(validate, 1, "graphical_interpretation_check is not 通过")
        write_rows(workspace / "validation_log.csv", [base_gate])

        page_only = dict(base_issue, drawing_refs="PDF第1页")
        write_rows(workspace / "issue_candidates.csv", [page_only])
        run(validate, 1, "drawing_refs missing linked fact drawing_no")

        leaked_filename = dict(base_issue, standard_display_name="匿名规范.pdf")
        write_rows(workspace / "issue_candidates.csv", [leaked_filename])
        run(validate, 1, "must not contain .pdf")

        fake_png = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00d\x00\x00\x00d"
        shot_a = workspace / "screenshots" / "a.png"
        shot_b = workspace / "screenshots" / "b.png"
        shot_a.write_bytes(fake_png)
        shot_b.write_bytes(fake_png)
        unexplained_multiple = dict(
            base_issue,
            needs_screenshot="yes",
            screenshot_strategy="multiple",
            screenshot_path="screenshots/a.png,screenshots/b.png",
            screenshot_location="A-01",
            screenshot_reason="",
            screenshot_count="2",
            evidence_point="匿名证据",
            red_box_target="匿名目标",
            context_required="保留图号和图名",
            screenshot_quality="清晰",
        )
        write_rows(workspace / "issue_candidates.csv", [unexplained_multiple])
        run(validate, 1, "multiple requires screenshot_reason")

        # v1.2 default report: internal-review title/style plus cited and no-citation display.
        second_check = check("C002", "I002")
        second_none = none_issue("I002", "C002", "2")
        write_rows(workspace / "check_matrix.csv", [base_check, second_check])
        write_rows(workspace / "issue_candidates.csv", [base_issue, second_none])
        write_rows(workspace / "validation_log.csv", [base_gate, gate("I002")])
        run(snapshot, 0)
        default_generate = run([
            python,
            str(SCRIPTS / "generate_review_report.py"),
            str(workspace),
            "--project-name",
            "匿名项目",
        ], 0)
        default_report = Path(default_generate.stdout.strip().splitlines()[-1])
        expected_name = f"【建单内审】匿名项目{date.today().isoformat()}.docx"
        if default_report.name != expected_name:
            raise AssertionError(f"unexpected v1.2 default report name: {default_report.name}")
        run([python, str(SCRIPTS / "validate_docx_content.py"), str(workspace), str(default_report)], 0)
        default_doc = Document(default_report)
        default_text = "\n".join(paragraph.text for paragraph in default_doc.paragraphs)
        expected_date = f"{date.today().year}年{date.today().month}月{date.today().day}日"
        for marker in ["建筑单体施工图内审意见", "设计说明：", expected_date, "《匿名规范》第1条", "【法规条文】：无。"]:
            if marker not in default_text:
                raise AssertionError(f"v1.2 default report missing: {marker}")
        if "一、设计说明：" in default_text or ".pdf" in default_text.casefold():
            raise AssertionError("v1.2 default report leaked a legacy heading or PDF filename")

        # v1.3: a reviewed sheet closes only after every required coverage topic is resolved.
        set_schema_version(workspace, "1.3")

        def v13_check(
            check_id: str,
            topic: str,
            evidence_class: str,
            actual_fact: str,
            *,
            conclusion: str = "符合",
            applicability: str = "适用",
            forms_issue: str = "no",
            issue_id: str = "",
        ) -> dict[str, str]:
            value = check(check_id, issue_id, forms_issue=forms_issue, conclusion=conclusion)
            value.update(
                review_item_id=f"ANON-{topic}",
                coverage_topic=topic,
                evidence_class=evidence_class,
                applicability=applicability,
                source_review_item=f"匿名审查要点：{topic}",
                actual_fact=actual_fact,
                fact_ids="F001",
                drawing_refs="A-01 一层平面图",
                not_forming_reason="" if forms_issue == "yes" else "已按主题核对并形成结论。",
            )
            return value

        v13_fact = dict(base_fact, raw_text_or_measure="匿名普通建筑平面事实。")
        v13_checks = [
            v13_check("C301", "function_layout", "relationship", "房间功能与相邻关系已逐项核对。"),
            v13_check("C302", "fire_egress", "numeric", "安全出口数量和疏散宽度已定量核对。"),
            v13_check("C303", "accessibility", "location", "无障碍通行位置与服务范围已核对。"),
            v13_check(
                "C304",
                "wet_room_hygiene",
                "not_applicable",
                "本图无有水房间，本主题不适用。",
                conclusion="不适用",
                applicability="不适用",
            ),
            v13_check("C305", "safety_coordination", "relationship", "平面安全边界及协调关系已核对。"),
        ]
        v13_inventory = dict(base_inventory, review_check_ids=";".join(item["check_id"] for item in v13_checks))
        write_rows(workspace / "drawing_inventory.csv", [v13_inventory])
        write_rows(workspace / "fact_ledger.csv", [v13_fact])
        write_rows(workspace / "check_matrix.csv", v13_checks)
        write_rows(workspace / "issue_candidates.csv", [])
        write_rows(workspace / "validation_log.csv", [])
        run(snapshot, 0)
        run(validate, 0)

        missing_topic_checks = [item for item in v13_checks if item["coverage_topic"] != "safety_coordination"]
        missing_topic_inventory = dict(
            v13_inventory,
            review_check_ids=";".join(item["check_id"] for item in missing_topic_checks),
        )
        write_rows(workspace / "drawing_inventory.csv", [missing_topic_inventory])
        write_rows(workspace / "check_matrix.csv", missing_topic_checks)
        run(validate, 1, "required coverage topics are not closed: safety_coordination")

        wall_inventory = dict(
            base_inventory,
            drawing_no="A-09",
            drawing_name="墙身大样图",
            content_type="大样图",
            review_family="墙身大样",
        )
        wall_fact = dict(
            base_fact,
            drawing_no="A-09",
            drawing_name="墙身大样图",
            raw_text_or_measure="节点标注300mm、500mm和600mm尺寸。",
        )
        wall_checks = [
            v13_check("C311", "detail_traceability", "relationship", "节点索引关系已核对。"),
            v13_check("C312", "window_sill_drainage", "numeric", "节点标注300mm控制尺寸。"),
            v13_check("C313", "parapet_flashing", "detail", "女儿墙泛水构造已核对。"),
            v13_check("C314", "waterproofing_detail", "detail", "防水收头构造已核对。"),
        ]
        for item in wall_checks:
            item["drawing_refs"] = "A-09 墙身大样图"
        wall_inventory["review_check_ids"] = ";".join(item["check_id"] for item in wall_checks)
        write_rows(workspace / "drawing_inventory.csv", [wall_inventory])
        write_rows(workspace / "fact_ledger.csv", [wall_fact])
        write_rows(workspace / "check_matrix.csv", wall_checks)
        run(validate, 1, "evidence_class numeric does not fit coverage_topic window_sill_drainage")

        dormitory_fact = dict(base_fact, raw_text_or_measure="本层为学生宿舍，设置电梯。")
        write_rows(workspace / "drawing_inventory.csv", [v13_inventory])
        write_rows(workspace / "fact_ledger.csv", [dormitory_fact])
        write_rows(workspace / "check_matrix.csv", v13_checks)
        run(validate, 1, "function-triggered coverage topics are not closed")

        dormitory_checks = v13_checks + [
            v13_check("C321", "dormitory_refuse", "location", "垃圾收集间位置和卫生措施已核对。"),
            v13_check("C322", "dormitory_acoustics", "performance", "宿舍与公共空间隔声减振措施已核对。"),
            v13_check("C323", "accessible_room", "numeric", "无障碍宿舍数量和位置已核对。"),
            v13_check("C324", "accessible_room_detail", "detail", "无障碍宿舍及卫生间详图已核对。"),
            v13_check("C325", "accessible_elevator", "performance", "无障碍电梯设施和提示要求已核对。"),
        ]
        dormitory_inventory = dict(
            v13_inventory,
            review_check_ids=";".join(item["check_id"] for item in dormitory_checks),
        )
        write_rows(workspace / "drawing_inventory.csv", [dormitory_inventory])
        write_rows(workspace / "check_matrix.csv", dormitory_checks)
        run(snapshot, 0)
        run(validate, 0)

        food_fact = dict(base_fact, raw_text_or_measure="本层设置食堂、餐厅和厨房。")
        food_check = v13_check(
            "C331",
            "food_wet_room_vertical",
            "numeric",
            "厨房开间标注6000mm。",
        )
        food_checks = v13_checks + [food_check]
        food_inventory = dict(
            v13_inventory,
            review_check_ids=";".join(item["check_id"] for item in food_checks),
        )
        write_rows(workspace / "drawing_inventory.csv", [food_inventory])
        write_rows(workspace / "fact_ledger.csv", [food_fact])
        write_rows(workspace / "check_matrix.csv", food_checks)
        run(validate, 1, "evidence_class numeric does not fit coverage_topic food_wet_room_vertical")

        design_inventory = dict(
            base_inventory,
            drawing_no="A-00",
            drawing_name="建筑设计说明",
            content_type="设计说明",
            review_family="设计说明",
        )
        design_fact = dict(
            base_fact,
            drawing_no="A-00",
            drawing_name="建筑设计说明",
            raw_text_or_measure="设计说明列出项目范围和设计依据。",
        )
        design_checks = [
            v13_check("C341", "project_scope_consistency", "relationship", "项目范围和楼层功能已核对。"),
            v13_check("C342", "code_basis_consistency", "identity_text", "设计依据名称及版本已核对。"),
        ]
        for item in design_checks:
            item["drawing_refs"] = "A-00 建筑设计说明"
        design_inventory["review_check_ids"] = ";".join(item["check_id"] for item in design_checks)
        write_rows(workspace / "drawing_inventory.csv", [design_inventory])
        write_rows(workspace / "fact_ledger.csv", [design_fact])
        write_rows(workspace / "check_matrix.csv", design_checks)
        run(validate, 1, "required coverage topics are not closed: project_identity_function")

        issue_checks = list(v13_checks)
        issue_checks[0] = v13_check(
            "C301",
            "function_layout",
            "relationship",
            "匿名功能关系与要求不一致。",
            conclusion="不符合",
            forms_issue="yes",
            issue_id="I301",
        )
        issue_inventory = dict(
            v13_inventory,
            review_check_ids=";".join(item["check_id"] for item in issue_checks),
        )
        first_v13_issue = issue("I301", "C301", "1", standard)
        duplicate_root_issue = issue("I302", "C301", "2", standard)
        write_rows(workspace / "drawing_inventory.csv", [issue_inventory])
        write_rows(workspace / "fact_ledger.csv", [v13_fact])
        write_rows(workspace / "check_matrix.csv", issue_checks)
        write_rows(workspace / "issue_candidates.csv", [first_v13_issue, duplicate_root_issue])
        write_rows(workspace / "validation_log.csv", [gate("I301"), gate("I302")])
        run(snapshot, 0)
        run(validate, 1, "verified issues share primary check_id: C301")

        mixed_primary_issue = dict(first_v13_issue, check_id="C301;C302")
        write_rows(workspace / "issue_candidates.csv", [mixed_primary_issue])
        write_rows(workspace / "validation_log.csv", [gate("I301")])
        run(snapshot, 0)
        run(validate, 1, "check_id not found in check_matrix.csv: C301;C302")

        write_rows(workspace / "issue_candidates.csv", [first_v13_issue])
        write_rows(workspace / "validation_log.csv", [gate("I301")])
        run(snapshot, 0)
        run(validate, 0)
        v13_generate = run([
            python,
            str(SCRIPTS / "generate_review_report.py"),
            str(workspace),
            "--project-name",
            "匿名项目",
        ], 0)
        v13_report = Path(v13_generate.stdout.strip().splitlines()[-1])
        run([python, str(SCRIPTS / "validate_docx_content.py"), str(workspace), str(v13_report)], 0)
        v13_text = "\n".join(paragraph.text for paragraph in Document(v13_report).paragraphs)
        if "建筑单体施工图内审意见" not in v13_text or "《匿名规范》第1条" not in v13_text:
            raise AssertionError("v1.3 report did not preserve the confirmed internal-review format")

    run([python, str(SKILL_DIR / "evals" / "test_v14_workflow.py")], 0, "12 v1.4 workflow")
    run([python, str(SKILL_DIR / "evals" / "test_review_calculations.py")], 0, "5 deterministic calculation")
    run([python, str(SKILL_DIR / "evals" / "test_cross_sheet_consistency.py")], 0, "4 cross-sheet consistency")
    print(f"PASS: {len(metadata)} anonymous regression scenarios completed")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
